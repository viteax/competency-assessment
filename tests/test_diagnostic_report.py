"""Отчет «Результаты проверки уровня сформированности компетенций» в Word (форма кафедры)."""

import io

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app import models
from app.services import diagnostic_report
from tests.helpers import answer_for, assign


def finish(admin: TestClient, student: TestClient, assignment_id: int, wrong: int) -> dict:
    """Студент сдает тест: на первые wrong вопросов отвечает неверно."""
    started = student.post("/me/attempts", json={"assignment_id": assignment_id}).json()
    bank = {q["id"]: q for q in admin.get("/questions").json()}
    answers = [
        {"question_id": q["id"], **answer_for(bank[q["id"]], correct=i >= wrong)}
        for i, q in enumerate(started["questions"])
    ]
    result = student.post(f"/me/attempts/{started['attempt_id']}/submit", json={"answers": answers})
    assert result.status_code == 200, result.text
    return started


def document_text(content: bytes) -> tuple[list[str], list[list[str]]]:
    doc = Document(io.BytesIO(content))
    table = doc.tables[0]
    return [p.text for p in doc.paragraphs], [[c.text for c in row.cells] for row in table.rows]


def test_report_for_group_test(admin: TestClient, login, student: TestClient):
    ctx = assign(admin, ["ПК-1", "ОПК-5"], per_competency=3)
    finish(admin, student, ctx["assignment"]["id"], wrong=0)
    second = login("student2", "student123")
    finish(admin, second, ctx["assignment"]["id"], wrong=6)  # все неверно

    response = admin.get(
        f"/reports/groups/{ctx['group']['id']}/results.docx",
        params={"test_id": ctx["test"]["id"]},
    )
    assert response.status_code == 200, response.text
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    paragraphs, rows = document_text(response.content)
    assert paragraphs[0] == "Результаты проверки уровня сформированности компетенций"
    assert "Группа: ПИ-21" in paragraphs
    assert (
        "Основная профессиональная образовательная программа: 09.03.04 Программная инженерия"
        in paragraphs
    )
    assert any(p.startswith("ОПК-5 ") for p in paragraphs) and any(
        p.startswith("ПК-1 ") for p in paragraphs
    )
    # ТФ направления, связанные с проверенными компетенциями.
    assert any("A/02.3" in p for p in paragraphs) and any("B/03.4" in p for p in paragraphs)
    assert paragraphs[-1] == (
        "Доля обучающихся, выполнивших 70% и более заданий диагностической работы, "
        "составляет 50.00 %"
    )

    assert rows[0][2] == "Задания закрытого типа" and rows[0][4] == "Задания открытого типа"
    assert rows[1][6] == "Номера неправильных ответов"
    data = rows[2:]
    assert len(data) == 2
    # Всего 6 заданий, у первого все верно, у второго все неверно.
    totals = sorted((r[7], r[8], r[9]) for r in data)
    assert totals == [("6", "0", "0"), ("6", "6", "100")]
    for r in data:
        closed, opened = int(r[2]), int(r[4])
        assert closed + opened == 6
    wrong_row = next(r for r in data if r[8] == "0")
    assert wrong_row[5] == "0" and len(wrong_row[6].split(",")) == int(wrong_row[4])


def test_report_numbers_follow_variant_document(admin: TestClient, student: TestClient, session):
    """Номера неверных открытых заданий – как в документе варианта: по компетенциям в
    порядке ФГОС (ОПК раньше ПК), внутри – по порядку в варианте, а не в перемешанном для
    студента порядке."""
    ctx = assign(admin, ["ПК-1", "ОПК-5"], per_competency=3)
    finish(admin, student, ctx["assignment"]["id"], wrong=6)
    report = diagnostic_report.collect(session, ctx["group"]["id"], ctx["test"]["id"])
    (row,) = report.rows
    attempt = session.scalars(select(models.Attempt)).one()
    variant = session.scalars(
        select(models.TestQuestion)
        .where(
            models.TestQuestion.test_id == ctx["test"]["id"],
            models.TestQuestion.variant == attempt.variant,
        )
        .order_by(models.TestQuestion.position)
    ).all()
    ordered = sorted(variant, key=lambda tq: tq.question.competency.code != "ОПК-5")
    expected = [i for i, tq in enumerate(ordered, start=1) if tq.question.kind.value == "open"]
    assert row.open_wrong == expected
    assert row.identifier == "student1"  # логин: номера в СДО у демо-студента нет


def test_report_requires_assignment(admin: TestClient):
    ctx = assign(admin, ["ПК-1"])
    other = admin.post("/groups", json={"name": "ПИ-25", "direction_id": 1}).json()
    response = admin.get(
        f"/reports/groups/{other['id']}/results.docx", params={"test_id": ctx["test"]["id"]}
    )
    assert response.status_code == 404 and "не назначен" in response.json()["detail"]
    missing = admin.get(f"/reports/groups/{ctx['group']['id']}/results.docx")
    assert missing.status_code == 422  # тест обязателен


def test_report_requires_admin(student: TestClient):
    assert student.get("/reports/groups/1/results.docx", params={"test_id": 1}).status_code == 403
