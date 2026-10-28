"""Вывод студенту по компетенциям и выгрузка «розы ветров» с подписью."""

import base64
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.services.reports import level_of
from tests.helpers import assign, correct_answers

# Минимальная PNG-картинка 1×1.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYGD4DwABBAEAwS2OUAAAAABJRU5ErkJggg=="
)


@pytest.mark.parametrize(
    ("ratio", "level"),
    [
        (0, "attention"),
        (0.39, "attention"),
        (0.4, "forming"),
        (0.79, "forming"),
        (0.8, "good"),
        (1, "good"),
    ],
)
def test_levels(ratio: float, level: str):
    assert level_of(ratio, 0.6) == level  # порог по умолчанию границ не двигает


@pytest.mark.parametrize(
    ("ratio", "threshold", "level"),
    [
        (0.8, 0.9, "forming"),  # ниже порога – не «хорошо развита»
        (0.9, 0.9, "good"),
        (0.3, 0.2, "forming"),  # порог ниже 40% – «внимание» только ниже порога
        (0.1, 0.2, "attention"),
    ],
)
def test_levels_follow_threshold(ratio: float, threshold: float, level: str):
    assert level_of(ratio, threshold) == level


def test_attempt_result_breaks_down_by_competency(admin: TestClient, student: TestClient):
    ctx = assign(admin, ["ОПК-1", "ПК-1"], per_competency=5)
    started = student.post("/me/attempts", json={"assignment_id": ctx["assignment"]["id"]}).json()
    correct = correct_answers(admin)
    bank = {q["id"]: q for q in admin.get("/questions").json()}
    competency = {c["id"]: c["code"] for c in admin.get("/competencies").json()}
    # ОПК-1 – все 5 верно, ПК-1 – только 1 из 5.
    answers, pk1 = [], 0
    for q in started["questions"]:
        code = competency[bank[q["id"]]["competency_id"]]
        if code == "ПК-1":
            pk1 += 1
            if pk1 > 1:
                continue
        answers.append({"question_id": q["id"], **correct[q["id"]]})
    result = student.post(
        f"/me/attempts/{started['attempt_id']}/submit", json={"answers": answers}
    ).json()

    assert (result["score"], result["max_score"]) == (6, 10)
    by_code = {c["code"]: c for c in result["competencies"]}
    assert list(by_code) == ["ОПК-1", "ПК-1"]
    assert (by_code["ОПК-1"]["correct"], by_code["ОПК-1"]["level"]) == (5, "good")
    assert (by_code["ПК-1"]["correct"], by_code["ПК-1"]["level"]) == (1, "attention")
    assert by_code["ПК-1"]["name"]  # формулировка компетенции для вывода

    report = student.get("/me/reports/competencies").json()["competencies"]
    assert [c["level"] for c in report] == ["good", "attention"]


def test_profile_legends(admin: TestClient, student: TestClient):
    ctx = assign(admin, ["ПК-1"])
    started = student.post("/me/attempts", json={"assignment_id": ctx["assignment"]["id"]}).json()
    student.post(f"/me/attempts/{started['attempt_id']}/submit", json={"answers": []})
    group = ctx["group"]

    mine = student.get("/me/reports/profile").json()["legend"]
    assert mine[0] == "Иванов Иван Иванович"
    assert mine[1].startswith("Группа ПИ-21 · ") and "курс · 09.03.04" in mine[1]

    legend = admin.get(f"/reports/groups/{group['id']}/profile").json()["legend"]
    assert legend[-1] == "Прошли тест: 1 из 3 студентов"
    test_id = ctx["test"]["id"]
    legend = admin.get(f"/reports/groups/{group['id']}/profile?test_id={test_id}").json()["legend"]
    assert legend[1] == f"Тест: {ctx['test']['title']}"


def test_chart_docx(student: TestClient):
    response = student.post(
        "/charts/docx",
        json={
            "title": "Компетенции – мои результаты",
            "legend": ["Иванов Иван Иванович", "Группа ПИ-21"],
            "rows": [["ПК-1", "Способен разрабатывать", "4 из 5"]],
            "image": "data:image/png;base64," + base64.b64encode(PNG).decode(),
        },
    )
    assert response.status_code == 200
    doc = Document(io.BytesIO(response.content))
    text = [p.text for p in doc.paragraphs]
    assert text[0] == "Компетенции – мои результаты"
    assert "Иванов Иван Иванович" in text and "Группа ПИ-21" in text
    assert doc.tables[0].rows[0].cells[2].text == "4 из 5"
    assert len(doc.inline_shapes) == 1


def test_chart_docx_rejects_non_png(student: TestClient, client: TestClient):
    body = {"title": "x", "image": base64.b64encode(b"<svg/>").decode()}
    assert student.post("/charts/docx", json=body).status_code == 400
    assert student.post("/charts/docx", json={**body, "image": "не base64"}).status_code == 400
    assert client.post("/charts/docx", json=body).status_code == 401
