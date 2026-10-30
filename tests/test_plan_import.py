"""Загрузка учебного плана из «Планов»: направление, компетенции, дисциплины с семестрами и
статусом, профстандарты, ОТФ, ТФ и связи ТФ ⇒ компетенции."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.services import plan_import
from tests.import_samples import plan_xlsx

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_parse_takes_everything_from_the_sheets():
    plan = plan_import.parse(plan_xlsx())
    assert (plan.direction_code, plan.direction_name, plan.year) == (
        "09.03.01",
        "Информатика и вычислительная техника",
        2023,
    )
    assert plan.profile == "Программное обеспечение робототехнических систем"
    # Формулировки – с прописной и без точки в конце.
    assert {c.code: c.description for c in plan.competencies.values()} == {
        "УК-1": "Способен осуществлять поиск и анализ информации",
        "УК-4": "Способен осуществлять деловую коммуникацию",
        "ОПК-1": "Способен применять естественнонаучные знания",
        "ПК-1": "Способен разрабатывать программное обеспечение",
    }
    # Модули и группы ДВ пропущены, ГИА – отдельно, семестры – по колонкам семестров (у
    # элективов по физкультуре – по часам), статус – по индексу.
    disciplines = {
        d.code: (d.name, d.kind.value, d.first_semester, d.last_semester, d.competencies)
        for d in plan.disciplines
    }
    assert disciplines == {
        "Б1.О.01.01": ("Иностранный язык", "regular", 1, 2, ["УК-4"]),
        "Б1.О.01.02": ("Математика", "regular", 1, 1, ["УК-1", "ОПК-1"]),
        "Б1.В.01.01": ("Базы данных", "regular", 3, 3, ["ОПК-1", "ПК-1"]),
        "Б1.В.01.02": ("Без семестров", "regular", None, None, ["ПК-1"]),
        "Б1.В.01.ДВ.01.01": ("Сети", "elective", 4, 4, ["ПК-1"]),
        "Б1.В.01.ДВ.01.02": ("Телекоммуникации", "elective", 4, 4, ["ПК-1"]),
        "Б1.В.02.ДВ.01.01": ("Игровые виды спорта", "elective", 2, 4, ["УК-1"]),
        "Б2.О.01(У)": ("Учебная практика", "practice", 2, 2, ["ОПК-1"]),
        "ФТД.01.01": ("Основы информационной безопасности", "optional", 3, 3, ["ПК-1", "ПК-2"]),
    }
    assert plan.skipped == ["Б3.01(Д) Выполнение и защита ВКР"]
    assert plan.warnings == [
        "Б1.В.01.02 «Без семестров»: в плане не найдены семестры – укажите их сами",
        "ПК-2 нет на листе «Компетенции» – связи с ней пропущены",
    ]
    tree = [
        (s.code, s.name, [(g.code, g.level, [(f.code, f.competencies) for f in g.functions])
                          for g in s.generals])
        for s in plan.standards
    ]  # fmt: skip
    assert tree == [
        ("06.001", "Программист", [("D", 6, [("D/03.6", ["ПК-1"])])]),
        ("06.022", "Системный аналитик", [("С", 6, [("С/01.6", ["ПК-1", "ОПК-1"])])]),
    ]


def upload(client: TestClient, path: str, data: bytes = b"", **form) -> dict:
    response = client.post(
        path, files={"file": ("plan.plx.xlsx", data or plan_xlsx(), XLSX)}, data=form
    )
    assert response.status_code == 200, response.text
    return response.json()


def by_code(items: list[dict]) -> dict[str, dict]:
    return {i["code"]: i for i in items}


def test_new_direction_from_plan(admin: TestClient):
    preview = upload(admin, "/import/plan/preview")
    assert preview["direction"]["id"] is None
    assert {c["status"] for c in preview["competencies"]} == {"new"}
    assert {d["status"] for d in preview["disciplines"]} == {"new"}
    standards = by_code(preview["standards"])
    assert standards["06.001"]["exists"] is True  # есть в демо-данных
    assert standards["06.001"]["general_functions"][0]["exists"] is False  # ОТФ D – новая
    assert standards["06.022"]["exists"] is False

    result = upload(admin, "/import/plan/apply")
    assert result == {
        "direction_id": result["direction_id"],
        "created_direction": True,
        "created_competencies": 4,
        "updated_competencies": 0,
        "created_disciplines": 9,
        "updated_disciplines": 0,
        "created_standards": 1,
        "created_general_functions": 2,
        "created_labor_functions": 2,
        "linked_labor_functions": 3,
    }
    direction = result["direction_id"]
    query = {"direction_id": direction}
    disciplines = {d["name"]: d for d in admin.get("/disciplines", params=query).json()}
    assert disciplines["Сети"]["kind"] == "elective"
    assert disciplines["Сети"]["code"] == "Б1.В.01.ДВ.01.01"
    assert disciplines["Учебная практика"]["kind"] == "practice"
    plan = {e["discipline_id"]: e for e in admin.get("/curriculum", params=query).json()}
    sport = plan[disciplines["Игровые виды спорта"]["id"]]
    assert (sport["first_semester"], sport["last_semester"]) == (2, 4)
    assert disciplines["Без семестров"]["id"] not in plan  # семестры не указаны – строки нет

    competencies = {c["id"]: c["code"] for c in admin.get("/competencies", params=query).json()}
    functions = {f["code"]: f for f in admin.get("/labor-functions").json()}
    assert sorted(competencies[i] for i in functions["С/01.6"]["competency_ids"]) == [
        "ОПК-1",
        "ПК-1",
    ]

    # Повторная загрузка того же плана ничего не меняет.
    again = upload(admin, "/import/plan/preview")
    assert again["direction"]["id"] == direction
    assert {i["status"] for i in again["competencies"] + again["disciplines"]} == {"same"}
    functions = [
        f for s in again["standards"] for g in s["general_functions"] for f in g["functions"]
    ]
    assert {f["status"] for f in functions} == {"same"}
    second = upload(admin, "/import/plan/apply")
    assert [v for k, v in second.items() if k != "direction_id"] == [False, 0, 0, 0, 0, 0, 0, 0, 0]


def test_plan_completes_existing_direction(admin: TestClient):
    direction = admin.post(
        "/directions", json={"code": "09.03.01", "name": "ИВТ (заведено вручную)"}
    ).json()
    uk1 = admin.post(
        "/competencies",
        json={
            "direction_id": direction["id"],
            "code": "УК-1",
            "kind": "УК",
            "description": "Старая",
        },
    ).json()
    math = admin.post(
        "/disciplines",
        json={"direction_id": direction["id"], "name": "Математика", "competency_ids": [uk1["id"]]},
    ).json()
    teacher = next(u for u in admin.get("/users").json() if u["role"] == "teacher")
    admin.post(
        "/curriculum",
        json={
            "discipline_id": math["id"],
            "first_semester": 2,
            "last_semester": 2,
            "teacher_id": teacher["id"],
            "quotas": [{"competency_id": uk1["id"], "questions": 7}],
        },
    )

    merge = {"merge_into": str(direction["id"])}
    preview = upload(admin, "/import/plan/preview", **merge)
    assert preview["direction"]["name"] == "ИВТ (заведено вручную)"
    item = by_code(preview["competencies"])["УК-1"]
    assert (item["status"], item["existing_description"]) == ("update", "Старая")
    assert by_code(preview["disciplines"])["Б1.О.01.02"]["changes"] == [
        "индекс – → Б1.О.01.02",
        "семестры 2 → 1",
        "+ ОПК-1",
    ]

    result = upload(admin, "/import/plan/apply", **merge)
    assert (result["created_direction"], result["updated_competencies"]) == (False, 0)
    assert result["updated_disciplines"] == 1
    query = {"direction_id": direction["id"]}
    assert by_code(admin.get("/competencies", params=query).json())["УК-1"]["description"] == (
        "Старая"
    )
    entry = next(
        e for e in admin.get("/curriculum", params=query).json() if e["discipline_id"] == math["id"]
    )
    # Семестры – из плана; преподаватель и квоты остаются.
    assert (entry["first_semester"], entry["teacher_id"], entry["quotas"]) == (
        1,
        teacher["id"],
        [{"competency_id": uk1["id"], "questions": 7}],
    )

    upload(admin, "/import/plan/apply", overwrite_descriptions="true")
    assert by_code(admin.get("/competencies", params=query).json())["УК-1"]["description"] == (
        "Способен осуществлять поиск и анализ информации"
    )


def test_labor_function_codes_match_in_both_alphabets(admin: TestClient):
    """В базе ОТФ «C» латиницей, в плане – «С» кириллицей: это одна ОТФ, дубля нет."""
    standard = admin.post(
        "/standards", json={"code": "06.022", "name": "Системный аналитик"}
    ).json()
    admin.post(
        "/general-functions",
        json={
            "standard_id": standard["id"],
            "code": "C",
            "name": "Проектирование",
            "qualification_level": 6,
        },
    )
    result = upload(admin, "/import/plan/apply")
    assert result["created_general_functions"] == 1  # только ОТФ D профстандарта 06.001
    generals = [
        g for g in admin.get("/general-functions").json() if g["standard_id"] == standard["id"]
    ]
    assert [(g["code"], g["name"]) for g in generals] == [("C", "Проектирование")]
    assert [f["code"] for f in generals[0]["functions"]] == ["С/01.6"]


def test_wrong_files(admin: TestClient):
    wb = Workbook()
    wb.active.append(["Код", "Наименование"])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin.post(
        "/import/plan/preview", files={"file": ("matrix.xlsx", buffer.getvalue(), XLSX)}
    )
    assert response.status_code == 400 and "не учебный план" in response.json()["detail"]
    response = admin.post("/import/plan/preview", files={"file": ("plan.docx", b"x", XLSX)})
    assert response.status_code == 400


def test_competency_import_sends_plan_to_its_tab(admin: TestClient):
    response = admin.post(
        "/import/preview",
        files={"file": ("plan.plx.xlsx", plan_xlsx(), XLSX)},
        data={"direction_id": "1"},
    )
    assert response.status_code == 400 and "«Загрузка УП»" in response.json()["detail"]


def test_xlsx_wording_ends_with_its_row(admin: TestClient):
    """Формулировка из Excel не захватывает следующие строки и листы (раньше последняя
    компетенция листа забирала текст всего файла)."""
    wb = Workbook()
    wb.active.append(["ПК-10", "Способен администрировать сети"])
    wb.active.append(["Итого часов", "1080", "Лек Лаб Пр"])
    wb.create_sheet("Другой лист").append(["№", "Индекс", "Наименование", "Семестр 1"])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = admin.post(
        "/import/preview",
        files={"file": ("matrix.xlsx", buffer.getvalue(), XLSX)},
        data={"direction_id": "1"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["competencies"][0]["description"] == "Способен администрировать сети"


@pytest.mark.parametrize("path", ["/import/plan/preview", "/import/plan/apply"])
def test_plan_import_requires_admin(student: TestClient, path: str):
    response = student.post(path, files={"file": ("plan.xlsx", plan_xlsx(), XLSX)})
    assert response.status_code == 403


# --- несколько учебных планов одного кода --------------------------------------------------

CODE = "09.03.01"


def _send(admin: TestClient, path: str, profile="Оптика", year="2023", **form):
    return admin.post(
        f"/import/plan/{path}",
        files={"file": ("plan.xlsx", plan_xlsx(profile, year), XLSX)},
        data=form,
    )


def _plans(admin: TestClient) -> list[dict]:
    return [d for d in admin.get("/directions").json() if d["code"] == CODE]


def _empty_plan(admin: TestClient, **fields) -> int:
    body = {"code": CODE, "name": "ИВТ", **fields}
    return admin.post("/directions", json=body).json()["id"]


def test_same_plan_is_updated(admin: TestClient):
    assert _send(admin, "apply").json()["created_direction"] is True
    assert _send(admin, "apply").json()["created_direction"] is False
    assert len(_plans(admin)) == 1
    assert _plans(admin)[0]["profile"] == "Оптика"
    assert _plans(admin)[0]["plan_year"] == 2023


def test_other_spelling_of_program_is_the_same_plan(admin: TestClient):
    _send(admin, "apply", profile="Оптика и связь")
    _send(admin, "apply", profile=" «ОПТИКА  и связь» ")
    assert len(_plans(admin)) == 1


def test_other_year_is_a_new_plan_without_questions(admin: TestClient):
    first = _send(admin, "apply", year="2023").json()["direction_id"]
    assert _send(admin, "preview", year="2026").json()["candidates"] == []
    second = _send(admin, "apply", year="2026").json()
    assert second["created_direction"] is True
    assert second["direction_id"] != first


def test_other_program_is_a_new_plan_without_questions(admin: TestClient):
    _send(admin, "apply", profile="Оптика")
    assert _send(admin, "preview", profile="Квантовая").json()["candidates"] == []
    assert _send(admin, "apply", profile="Квантовая").status_code == 200
    assert len(_plans(admin)) == 2


def test_plan_with_empty_fields_requires_a_choice(admin: TestClient):
    plan = _empty_plan(admin)
    candidates = _send(admin, "preview").json()["candidates"]
    assert candidates == [{"id": plan, "label": f"{CODE} ИВТ"}]
    refused = _send(admin, "apply")
    assert refused.status_code == 400
    assert "выберите" in refused.json()["detail"]
    assert len(_plans(admin)) == 1


def test_merge_fills_the_empty_fields(admin: TestClient):
    plan = _empty_plan(admin)
    result = _send(admin, "apply", merge_into=str(plan)).json()
    assert result["direction_id"] == plan
    assert result["created_direction"] is False
    [found] = _plans(admin)
    assert (found["profile"], found["plan_year"]) == ("Оптика", 2023)


def test_merge_preview_compares_with_chosen_plan(admin: TestClient):
    plan = _empty_plan(admin)
    uk1 = {"direction_id": plan, "code": "УК-1", "kind": "УК", "description": "Иное"}
    assert admin.post("/competencies", json=uk1).status_code == 201
    preview = _send(admin, "preview", merge_into=str(plan)).json()
    assert preview["direction"]["id"] == plan
    statuses = {c["code"]: c["status"] for c in preview["competencies"]}
    assert statuses["УК-1"] == "update"


def test_create_new_keeps_the_empty_plan(admin: TestClient):
    plan = _empty_plan(admin)
    result = _send(admin, "apply", create_new="true").json()
    assert result["created_direction"] is True
    assert {p["id"] for p in _plans(admin)} == {plan, result["direction_id"]}
    assert next(p for p in _plans(admin) if p["id"] == plan)["profile"] is None


def test_merge_into_a_conflicting_plan_is_refused(admin: TestClient):
    plan = _empty_plan(admin, profile="Другая")
    refused = _send(admin, "apply", merge_into=str(plan))
    assert refused.status_code == 400
    assert next(p for p in _plans(admin) if p["id"] == plan)["profile"] == "Другая"


def test_partial_match_is_a_candidate(admin: TestClient):
    plan = _empty_plan(admin, profile="Оптика")
    assert [c["id"] for c in _send(admin, "preview").json()["candidates"]] == [plan]


def test_file_without_program_and_year_warns_and_loads(admin: TestClient):
    preview = _send(admin, "preview", profile=None, year=None).json()
    assert any("программ" in w for w in preview["warnings"])
    assert any("год" in w for w in preview["warnings"])
    assert _send(admin, "apply", profile=None, year=None).status_code == 200
    [found] = _plans(admin)
    assert (found["profile"], found["plan_year"]) == (None, None)
