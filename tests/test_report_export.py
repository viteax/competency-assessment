"""Отчётность: фильтр по тесту, выгрузка в Excel, данные для «розы ветров»."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import answer_for, assign, student_id


def take(student: TestClient, admin: TestClient, assignment_id: int, right: int) -> None:
    """Студент проходит тест, отвечая верно на первые `right` вопросов своего варианта."""
    body = student.post("/me/attempts", json={"assignment_id": assignment_id}).json()
    questions = {q["id"]: q for q in admin.get("/questions").json()}
    answers = [
        {"question_id": q["id"], **answer_for(questions[q["id"]], correct=i < right)}
        for i, q in enumerate(body["questions"])
    ]
    result = student.post(f"/me/attempts/{body['attempt_id']}/submit", json={"answers": answers})
    assert result.status_code == 200, result.text


def test_group_report_filters_by_test_and_has_date(admin: TestClient, student: TestClient):
    first = assign(admin, ["ОПК-5"])
    take(student, admin, first["assignment"]["id"], right=3)

    group_id = first["group"]["id"]
    report = admin.get(f"/reports/groups/{group_id}/competencies").json()
    ivanov = next(s for s in report["students"] if s["full_name"] == "Иванов Иван Иванович")
    assert ivanov["tested_at"] is not None
    assert [(c["code"], c["correct"], c["total"]) for c in ivanov["competencies"]] == [
        ("ОПК-5", 3, 3)
    ]
    others = [s for s in report["students"] if s["student_id"] != ivanov["student_id"]]
    assert all(s["tested_at"] is None and not s["competencies"] for s in others)

    by_test = admin.get(
        f"/reports/groups/{group_id}/competencies", params={"test_id": first["test"]["id"]}
    ).json()
    assert by_test == report
    assert (
        admin.get(f"/reports/groups/{group_id}/competencies", params={"test_id": 999}).status_code
        == 404
    )


def test_other_test_results_are_excluded(admin: TestClient, student: TestClient):
    a = assign(admin, ["ОПК-5"])
    take(student, admin, a["assignment"]["id"], right=3)
    group = a["group"]
    other = admin.post(
        "/tests",
        json={
            "title": "Другой",
            "competency_ids": [
                c["id"] for c in admin.get("/competencies").json() if c["code"] == "ОПК-1"
            ],
            "questions_per_competency": 2,
        },
    ).json()
    admin.post("/assignments", json={"test_id": other["id"], "group_id": group["id"]})
    report = admin.get(
        f"/reports/groups/{group['id']}/competencies", params={"test_id": other["id"]}
    ).json()
    assert all(not s["competencies"] for s in report["students"])


def test_excel_export(admin: TestClient, student: TestClient):
    a = assign(admin, ["ОПК-5", "ПК-1"])
    take(student, admin, a["assignment"]["id"], right=4)  # 6 вопросов: 3 + 3
    response = admin.get(f"/reports/groups/{a['group']['id']}/competencies.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Компетенции", "Трудовые функции"]
    ws = wb["Компетенции"]
    header = [c.value for c in ws[1]]
    assert header == ["ФИО", "Группа", "Дата тестирования", "ОПК-5", "ПК-1"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    ivanov = rows["Иванов Иван Иванович"]
    assert ivanov[1] == "ПИ-21" and ivanov[2]
    # Порядок вопросов перемешан: сумма верных — 4 из 6, по компетенциям — «N из 3».
    cells = ivanov[3:5]
    assert all(c.endswith(" из 3") for c in cells)
    assert sum(int(c.split()[0]) for c in cells) == 4
    assert rows["Петрова Анна Сергеевна"][3:5] == (None, None)

    tf = wb["Трудовые функции"]
    tf_header = [c.value for c in tf[1]]
    assert tf_header[:2] == ["ФИО", "Группа"]
    assert any(h.startswith("A/02.3") for h in tf_header[2:])
    tf_rows = {r[0]: r for r in tf.iter_rows(min_row=2, values_only=True) if r[0]}
    a023 = tf_header.index(next(h for h in tf_header if h and h.startswith("A/02.3")))
    assert tf_rows["Иванов Иван Иванович"][a023].endswith("%")  # нужен только ПК-1
    b023 = tf_header.index(next(h for h in tf_header if h and h.startswith("B/02.4")))
    assert tf_rows["Иванов Иван Иванович"][b023] in (None, "")  # ПК-2 не проверялась


def test_profile_for_student_and_group(admin: TestClient, login):
    a = assign(admin, ["ОПК-5"])
    take(login("student1", "student123"), admin, a["assignment"]["id"], right=3)
    take(login("student2", "student123"), admin, a["assignment"]["id"], right=0)
    group_id = a["group"]["id"]

    group = admin.get(f"/reports/groups/{group_id}/profile").json()
    assert group["students"] == 2
    [opk5] = group["competencies"]
    assert (opk5["code"], opk5["value"], opk5["measured"], opk5["total"]) == ("ОПК-5", 0.5, 2, 3)
    tf = {t["code"]: t for t in group["labor_functions"]}
    assert tf["B/03.4"]["value"] == 0.5  # ПК-2 + ОПК-5, проверена только ОПК-5
    assert (tf["B/03.4"]["measured"], tf["B/03.4"]["total"]) == (1, 2)
    assert tf["A/02.3"]["value"] is None  # ПК-1 не проверялась

    ivanov = admin.get(
        f"/reports/groups/{group_id}/profile", params={"student_id": student_id(admin)}
    ).json()
    assert ivanov["competencies"][0]["value"] == 1.0
    stranger = admin.get(f"/reports/groups/{group_id}/profile", params={"student_id": 999})
    assert stranger.status_code == 404


def test_student_sees_own_profile(admin: TestClient, student: TestClient):
    a = assign(admin, ["ОПК-5"])
    take(student, admin, a["assignment"]["id"], right=3)
    mine = student.get("/me/reports/profile").json()
    assert [(c["code"], c["value"]) for c in mine["competencies"]] == [("ОПК-5", 1.0)]
    assert student.get(f"/reports/groups/{a['group']['id']}/profile").status_code == 403
