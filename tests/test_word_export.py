"""Выгрузка в Word по образцу ФОМ: тест по вариантам и ФОМ дисциплины, бланк и ключ."""

import io
import re

from docx import Document
from fastapi.testclient import TestClient

from tests.helpers import assign, by_code, discipline_id

TASK = re.compile(r"^\d+\. Задание ")


def load(response) -> Document:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return Document(io.BytesIO(response.content))


def texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def header_values(doc: Document) -> dict[str, str]:
    return {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}


def answer_grids(doc: Document) -> list[list[str]]:
    """Последние строки таблиц «Ответ:» (соответствие и последовательность)."""
    grids, body = [], doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    previous = None
    for child in body.iterchildren():
        if child in tables and previous is not None and previous.text.startswith("Ответ:"):
            grids.append([c.text for c in tables[child].rows[-1].cells])
        previous = paragraphs.get(child)
    return grids


TITLES = {
    "matching": "Задание на установление соответствия",
    "sequence": "Задание на установление последовательности",
    "open": "Задание открытого типа",
}


def answer_grid_after(doc: Document, question_text: str) -> list[str]:
    """Таблица «Ответ:» задания с этим текстом (соответствие или последовательность)."""
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    seen, previous = False, None
    for child in doc.element.body.iterchildren():
        paragraph = paragraphs.get(child)
        seen = seen or (paragraph is not None and paragraph.text == question_text)
        if seen and child in tables and previous is not None and previous.text.startswith("Ответ:"):
            return [c.text for c in tables[child].rows[-1].cells]
        previous = paragraph
    raise AssertionError(f"нет таблицы ответа после «{question_text}»")


def test_discipline_fund_blank_and_key(admin: TestClient):
    os_id = discipline_id(admin, "Операционные системы")
    bank = sorted(
        (q for q in admin.get("/questions").json() if q["discipline_id"] == os_id),
        key=lambda q: q["id"],
    )
    blank = load(admin.get(f"/disciplines/{os_id}/fund.docx"))
    assert header_values(blank) == {
        "Направления (специальности)": "09.03.04 Программная инженерия",
        "Семестр": "3",
        "Дисциплина": "Операционные системы",
        "Код": "ОПК-5.",
        "Наименование компетенции": f"ОПК-5 {by_code(admin, 'ОПК-5')['description']}",
    }
    lines = texts(blank)
    headings = [t for t in lines if TASK.match(t)]
    assert headings == [f"{i}. {TITLES[q['kind']]}" for i, q in enumerate(bank, start=1)]
    assert {q["kind"] for q in bank} == set(TITLES)  # в демо-ФОМ есть все три типа
    assert "Распределение тестовых заданий по видам и уровням сложности" in lines
    assert all(cell == "" for grid in answer_grids(blank) for cell in grid)
    assert "Ответ: ________________." in lines

    key = load(admin.get(f"/disciplines/{os_id}/fund.docx", params={"key": True}))
    grids = answer_grids(key)
    assert len(grids) == sum(q["kind"] != "open" for q in bank)
    assert all(cell.isdigit() for grid in grids for cell in grid)
    assert "Ответ: root." in texts(key)

    distribution = key.tables[-1]
    assert [c.text for c in distribution.rows[0].cells] == [
        "Семестр",
        "Код компетенции",
        "Дисциплина",
        "№ задания",
        "Виды задания",
        "Уровень сложности",
        "Время выполнения",
    ]
    nginx = next(i for i, q in enumerate(bank, start=1) if "nginx в Ubuntu" in q["text"])
    assert [c.text for c in distribution.rows[nginx].cells[3:]] == [
        str(nginx),
        "Задание на установление последовательности",
        "повышенный",
        "3-5",
    ]


def test_sequence_key_matches_blank_numbering(admin: TestClient):
    """В ключе – номера элементов бланка в правильном порядке."""
    os_id = discipline_id(admin, "Операционные системы")
    key = load(admin.get(f"/disciplines/{os_id}/fund.docx", params={"key": True}))
    lines = texts(key)
    question = (
        "Расположите команды установки и запуска веб-сервера nginx в Ubuntu в правильном порядке:"
    )
    start = lines.index(question)
    shown = {}
    for line in lines[start + 1 : start + 8]:
        if line[:1].isdigit():
            number, text = line.split(". ", 1)
            shown[text] = number
    correct = [
        "sudo apt update",
        "sudo apt install nginx",
        "sudo systemctl enable --now nginx",
        "systemctl status nginx",
    ]
    sequence_grid = answer_grid_after(key, question)
    assert sequence_grid == [shown[t] for t in correct]
    assert sequence_grid != ["1", "2", "3", "4"]  # показ не совпадает с ответом


def test_test_document_by_variants(admin: TestClient):
    ctx = assign(admin, ["ОПК-5"], per_competency=5)  # 4 варианта
    test_id = ctx["test"]["id"]
    whole = load(admin.get(f"/tests/{test_id}/document.docx"))
    titles = [t for t in texts(whole) if t.startswith("«Входной контроль», вариант")]
    assert titles == [f"«Входной контроль», вариант {n}" for n in range(1, 5)]

    one = admin.get(f"/tests/{test_id}/document.docx", params={"variant": 2, "key": True})
    assert "%D0%B2%D0%B0%D1%80%D0%B8%D0%B0%D0%BD%D1%82%202" in one.headers["content-disposition"]
    doc = load(one)
    assert [t for t in texts(doc) if t.startswith("«")] == ["«Входной контроль», вариант 2"]
    numbered = [t for t in texts(doc) if TASK.match(t)]
    assert len(numbered) == 5

    assert admin.get(f"/tests/{test_id}/document.docx", params={"variant": 9}).status_code == 404


def test_errors_and_access(admin: TestClient, student: TestClient):
    empty = admin.post("/disciplines", json={"direction_id": 1, "name": "Пустая"}).json()
    assert admin.get(f"/disciplines/{empty['id']}/fund.docx").status_code == 404
    assert admin.get("/tests/999/document.docx").status_code == 404
    os_id = discipline_id(admin, "Операционные системы")
    assert student.get(f"/disciplines/{os_id}/fund.docx").status_code == 403
