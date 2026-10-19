"""Выгрузка в Word из интерфейса: ФОМ дисциплины и собранный тест."""

import io

import pytest
from docx import Document
from playwright.sync_api import expect

from tests.e2e.conftest import assign_new_test, login

pytestmark = pytest.mark.e2e


def downloaded(page, link_text: str) -> tuple[str, Document]:
    with page.expect_download() as download:
        page.click(f"a:text-is('{link_text}')")
    content = open(download.value.path(), "rb").read()
    return download.value.suggested_filename, Document(io.BytesIO(content))


def test_fund_and_test_documents(new_page):
    page = new_page(accept_downloads=True)
    login(page, "admin", "admin123")
    page.click("nav >> text=Банк вопросов")
    expect(page.locator("text=Выберите дисциплину, чтобы выгрузить")).to_be_visible()
    page.get_by_label("Дисциплина", exact=True).select_option(label="Операционные системы")
    name, doc = downloaded(page, "ФОМ дисциплины (Word)")
    assert name == "ФОМ — Операционные системы.docx"
    assert doc.paragraphs[0].text == "Фонд оценочных материалов дисциплины"
    name, _ = downloaded(page, "ФОМ с ответами (Word)")
    assert name == "ФОМ — Операционные системы — ключ.docx"

    page.click("nav >> text=Назначить тест")
    assign_new_test(page, "Контрольная", ["ОПК-5"], per_competency=5)
    page.click("nav >> text=Тесты")
    page.click("text=Открыть")
    page.get_by_label("Вариант для Word").select_option(label="Вариант 3")
    name, doc = downloaded(page, "Тест с ответами (Word)")
    assert name == "Контрольная, вариант 3 — ключ.docx"
    assert "«Контрольная», вариант 3" in [p.text for p in doc.paragraphs]
