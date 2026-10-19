"""Отчётность: выгрузка в Excel и «розы ветров» у администратора и студента."""

import io
import re

import pytest
from docx import Document
from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from tests.e2e.conftest import answer_all, assign_new_test, login, logout

pytestmark = pytest.mark.e2e


def pass_test(page: Page, user: str) -> None:
    login(page, user, "student123")
    page.click("text=Пройти")
    questions = page.locator(".question")
    expect(questions.first).to_be_visible()
    answer_all(page)
    page.click("text=Сдать тест")
    expect(page.locator(".score")).to_be_visible()


def test_excel_and_radars(new_page):
    page = new_page(accept_downloads=True)
    login(page, "admin", "admin123")
    assign_new_test(page, "Остаточные знания", ["ОПК-1", "ОПК-5", "ПК-1"], per_competency=2)
    logout(page)
    pass_test(page, "student1")

    # Студент видит свои «розы».
    page.click("text=Мои результаты")
    expect(page.locator("h2:has-text('Компетенции – мои результаты')")).to_be_visible()
    expect(page.locator(".radar-wrap svg").first).to_be_visible()
    expect(page.locator(".radar-legend").first).to_contain_text("Группа ПИ-21")
    expect(page.locator("h2:text-is('Вывод по компетенциям')")).to_be_visible()
    logout(page)

    login(page, "admin", "admin123")
    page.click("nav >> text=Результаты группы")
    expect(page.locator("td:has-text('Иванов Иван Иванович')")).to_be_visible()
    page.get_by_label("Тест").select_option(label="Остаточные знания")
    expect(page.locator("th:text-is('ОПК-5')")).to_be_visible()

    with page.expect_download() as download:
        page.click("text=Скачать Excel")
    assert download.value.suggested_filename == "Результаты ПИ-21.xlsx"
    wb = load_workbook(io.BytesIO(open(download.value.path(), "rb").read()))
    header = [c.value for c in wb["Компетенции"][1]]
    assert header == ["ФИО", "Группа", "Дата тестирования", "ОПК-1", "ОПК-5", "ПК-1"]

    # Три компетенции — лепестковая диаграмма для группы и для студента.
    expect(page.locator("h2:has-text('Компетенции – группа ПИ-21, среднее')")).to_be_visible()
    competencies = page.locator(".radar-card").first
    expect(competencies.locator("svg text", has_text=re.compile("^ПК-1$"))).to_be_visible()
    expect(competencies.locator(".radar-legend")).to_contain_text("Прошли тест: 1 из 3 студентов")
    competencies.locator("circle[tabindex]").last.hover()  # при нулях точки совпадают
    expect(competencies.locator(".radar-tip")).to_contain_text("из 5")
    page.get_by_label("Чья роза").select_option(label="Иванов Иван Иванович")
    expect(page.locator("h2:has-text('Компетенции – Иванов Иван Иванович')")).to_be_visible()

    competencies = page.locator(".radar-card").first
    expect(competencies.locator(".radar-legend")).to_contain_text("Иванов Иван Иванович")
    with page.expect_download() as png:
        competencies.locator("button:text-is('PNG')").click()
    assert png.value.suggested_filename == "Компетенции – Иванов Иван Иванович.png"
    assert open(png.value.path(), "rb").read(8) == b"\x89PNG\r\n\x1a\n"
    with page.expect_download() as jpeg:
        competencies.locator("button:text-is('JPEG')").click()
    assert jpeg.value.suggested_filename.endswith(".jpg")
    assert open(jpeg.value.path(), "rb").read(3) == b"\xff\xd8\xff"
    with page.expect_download() as word:
        competencies.locator("button:text-is('Word')").click()
    assert word.value.suggested_filename == "Компетенции – Иванов Иван Иванович.docx"
    text = "\n".join(p.text for p in Document(word.value.path()).paragraphs)
    assert "Иванов Иван Иванович" in text and "Группа ПИ-21" in text


def test_too_few_axes_explained(admin_page: Page, new_page):
    page = admin_page
    assign_new_test(page, "Один ОПК", ["ОПК-5"], per_competency=2)
    student = new_page()
    pass_test(student, "student1")
    page.click("nav >> text=Результаты группы")
    # Одна компетенция и одна ТФ (B/03.4) — для обеих «роз» объяснение вместо диаграммы.
    expect(page.locator("text=Для диаграммы нужно минимум 3 оси, сейчас 1")).to_have_count(2)


def test_radar_section_explains_when_empty(admin_page: Page):
    page = admin_page
    page.click("nav >> text=Результаты группы")
    expect(page.locator(".notice")).to_contain_text("Появится, когда студенты сдадут тест")
    page.click("button:text-is('Роза ветров ↓')")
    expect(page.locator("h2:text-is('Роза ветров')")).to_be_in_viewport()
