"""Загрузка вопросов из фонда оценочных материалов (ФОС, ФОМ) дисциплины в Word.

Документ – как у кафедры и как выгрузка ФОМ из банка вопросов: таблица «Информация о
дисциплине» (направление, семестр, дисциплина, коды и формулировки компетенций), затем по
каждой компетенции – абзац с ее кодом и задания. У задания: заголовок «N. Задание на
установление соответствия / последовательности / открытого типа», таблица «Сценарий
выполнения – Результат оценивания», текст задания и «Ответ:» – текстом или таблицей (буквы и
номера – для соответствия, номера – для последовательности). Элементы списков бывают
пронумерованы текстом («1. …») или автонумерацией Word; маркированные подпункты относятся к
предыдущему элементу. В конце – таблица «Распределение заданий по видам и уровням
сложности»: из нее берется уровень сложности.

Можно загрузить один .docx или .zip с несколькими. Сначала предпросмотр, потом сохранение;
вопросы, которые уже есть в банке (тот же текст у той же дисциплины и компетенции), не
дублируются.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath

from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app import models, schemas
from app.errors import DomainError
from app.services import files, questions, teaching
from app.services.common import code_key, get_or_404

LETTERS = "АБВГДЕЖЗИКЛМНОПРСТУФХЦЧШЭЮЯ"
MAX_FILES = 50  # документов в архиве
MAX_UNPACKED = 50 * 1024 * 1024  # байт в распакованном архиве

_KINDS = {
    "задание на установление соответствия": models.QuestionKind.MATCHING,
    "задание на установление последовательности": models.QuestionKind.SEQUENCE,
    "задание открытого типа": models.QuestionKind.OPEN,
}
_DIFFICULTY = {"базовый": models.Difficulty.BASIC, "повышенный": models.Difficulty.ADVANCED}

_NUMBER = re.compile(r"^(\d{1,3})\s*[.)]\s*")
_LETTER = re.compile(r"^([А-ЯЁA-Z])\s*[.)]\s*")
_COMPETENCY = re.compile(r"^(УК|ОПК|ПК)\s?[-‑–—]\s?(\d+)(?![.,]?\d)[.:]?\s*(.*)$", re.S)
_DIRECTION = re.compile(r"\d{2}\.\d{2}\.\d{2}")
_BLANK = re.compile(r"^[_\s.]*$")  # «Ответ: ________.» в бланке – ответа нет
_ANSWER = re.compile(r"(?:^|\s)Ответ\s*[:–-]\s*(.*)$", re.S)


# --- блоки документа ------------------------------------------------------------------


@dataclass
class Para:
    text: str
    numbered: str | None = None  # автонумерация Word: "decimal", "bullet", … или None


Cell = list[Para]


@dataclass
class Grid:
    rows: list[list[Cell]]  # продолжение объединенной по вертикали ячейки – пустое

    def texts(self, row: int) -> list[str]:
        return [" ".join(p.text for p in cell if p.text) for cell in self.rows[row]]


class _Numbering:
    """Формат автонумерации абзаца ("decimal", "bullet", "russianUpper", …): из свойств абзаца
    или из его стиля («Список нумерованный» и т. п.)."""

    def __init__(self, document) -> None:
        self.formats: dict[tuple[str, str], str] = {}  # (numId, уровень) → формат
        self.styles: dict[str, tuple[str, str]] = {}  # стиль → (numId, уровень)
        try:
            numbering = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return
        levels = {}
        for abstract in numbering.findall(qn("w:abstractNum")):
            for lvl in abstract.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                key = (abstract.get(qn("w:abstractNumId")), lvl.get(qn("w:ilvl")))
                levels[key] = fmt.get(qn("w:val")) if fmt is not None else "decimal"
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            for (abstract_id, level), fmt in levels.items():
                if ref is not None and abstract_id == ref.get(qn("w:val")):
                    self.formats[(num.get(qn("w:numId")), level)] = fmt
        based_on = {}
        for style in document.styles.element.findall(qn("w:style")):
            style_id = style.get(qn("w:styleId"))
            num = self._num(style)
            if num is not None:
                self.styles[style_id] = num
            parent = style.find(qn("w:basedOn"))
            if parent is not None:
                based_on[style_id] = parent.get(qn("w:val"))
        for style_id, parent in based_on.items():
            seen = {style_id}
            while style_id not in self.styles and parent in based_on and parent not in seen:
                seen.add(parent)
                parent = based_on[parent]
            if style_id not in self.styles and parent in self.styles:
                self.styles[style_id] = self.styles[parent]

    @staticmethod
    def _num(element) -> tuple[str, str] | None:
        num = element.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        num_id = num.find(qn("w:numId")) if num is not None else None
        if num_id is None:
            return None
        level = num.find(qn("w:ilvl"))
        return num_id.get(qn("w:val")), level.get(qn("w:val")) if level is not None else "0"

    def of(self, p) -> str | None:
        num = self._num(p)
        if num is None:
            style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
            num = self.styles.get(style.get(qn("w:val"))) if style is not None else None
        if num is None or num[0] == "0":  # numId 0 – нумерация снята
            return None
        return self.formats.get(num, "decimal")


def _para(p, numbering: _Numbering) -> list[Para]:
    """Абзац; разрыв строки внутри абзаца – как начало нового."""
    lines, chunk = [], []
    for node in p.iter():
        if node.tag == qn("w:t"):
            chunk.append(node.text or "")
        elif node.tag == qn("w:tab"):
            chunk.append(" ")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            lines.append("".join(chunk))
            chunk = []
    lines.append("".join(chunk))
    numbered = numbering.of(p)
    return [
        Para(" ".join(line.split()), numbered if i == 0 else None) for i, line in enumerate(lines)
    ]


def _blocks(data: bytes) -> list[Para | Grid]:
    document = Document(io.BytesIO(data))
    numbering = _Numbering(document)
    blocks: list[Para | Grid] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.extend(_para(child, numbering))
        elif child.tag == qn("w:tbl"):
            rows = [
                [
                    [p for el in tc.findall(qn("w:p")) for p in _para(el, numbering)]
                    for tc in tr.findall(qn("w:tc"))
                ]
                for tr in child.findall(qn("w:tr"))
            ]
            blocks.append(Grid(rows))
    return blocks


def _norm(text: str) -> str:
    return " ".join(text.split()).casefold().replace("ё", "е")


# --- разбор ---------------------------------------------------------------------------


@dataclass
class FundTask:
    number: int  # по порядку в разделе компетенции
    competency: str | None  # код из заголовка раздела
    kind: models.QuestionKind
    text: str = ""
    options: list[schemas.AnswerOptionIn] = field(default_factory=list)
    difficulty: models.Difficulty = models.Difficulty.BASIC
    error: str | None = None


@dataclass
class FundDocument:
    filename: str
    direction: str  # как в шапке: «15.04.06 Мехатроника и робототехника …»
    semester: int | None
    discipline: str | None
    competencies: dict[str, str]  # код → формулировка из шапки
    tasks: list[FundTask]
    warnings: list[str]
    error: str | None = None  # документ не разобран (не ФОС, поврежден)

    @property
    def direction_code(self) -> str | None:
        m = _DIRECTION.search(self.direction)
        return m.group() if m else None


def _code(kind: str, number: str) -> str:
    return f"{kind}-{int(number)}"


def _header(grid: Grid) -> tuple[str, int | None, str | None, dict[str, str]] | None:
    """Шапка «Информация о дисциплине»: направление, семестр, дисциплина, компетенции."""
    labels = [_norm(grid.texts(r)[0]) for r in range(len(grid.rows)) if grid.rows[r]]
    if "дисциплина" not in labels:
        return None
    direction, semester, discipline = "", None, None
    competencies: dict[str, str] = {}
    pending: list[str] = []  # коды из строки «Код», ждущие формулировок
    for r, row in enumerate(grid.rows):
        if len(row) < 2:
            continue
        label = _norm(grid.texts(r)[0])
        value_paras = [p.text for p in row[1] if p.text]
        value = " ".join(value_paras)
        if label.startswith("направлени"):
            direction = value
        elif label == "семестр":
            digits = re.findall(r"\d+", value)
            semester = int(digits[0]) if digits else None
        elif label == "дисциплина":
            discipline = value or None
        elif label.startswith("код"):
            pending = [_code(k, n) for k, n in re.findall(r"(УК|ОПК|ПК)\s?[-‑–]\s?(\d+)", value)]
            for code in pending:
                competencies.setdefault(code, "")
        elif label.startswith("наименование"):
            # Формулировки: по строке на компетенцию («ОПК-9 Способен …») или одна – к
            # единственному коду из строки «Код».
            for line in value_paras:
                m = _COMPETENCY.match(line)
                if m:
                    competencies[_code(m[1], m[2])] = m[3].strip(" .")
                elif len(pending) == 1:
                    competencies[pending[0]] = line.strip(" .")
    return direction, semester, discipline, competencies


def _distribution(grid: Grid) -> list[tuple[str | None, int, models.Difficulty | None]] | None:
    """Таблица «Распределение заданий»: (код компетенции, № задания, уровень сложности)."""
    head = [_norm(t) for t in grid.texts(0)]
    number = next((i for i, h in enumerate(head) if "№" in h and "задани" in h), None)
    level = next((i for i, h in enumerate(head) if "сложност" in h), None)
    if number is None or level is None:
        return None
    code_col = next((i for i, h in enumerate(head) if "компетенц" in h), None)
    rows, code = [], None
    for r in range(1, len(grid.rows)):
        texts = grid.texts(r)
        if code_col is not None and code_col < len(texts):
            m = _COMPETENCY.match(texts[code_col])
            code = _code(m[1], m[2]) if m else code
        digits = re.findall(r"\d+", texts[number]) if number < len(texts) else []
        if digits:
            difficulty = _DIFFICULTY.get(_norm(texts[level])) if level < len(texts) else None
            rows.append((code, int(digits[0]), difficulty))
    return rows


def _task_kind(text: str) -> models.QuestionKind | None:
    norm = _norm(_NUMBER.sub("", text))
    return next((kind for title, kind in _KINDS.items() if norm.startswith(title)), None)


def _items(paras: list[Para]) -> list[tuple[int, str]]:
    """Нумерованные элементы списка: (номер, текст). Номер – из текста («2. …») или по порядку
    при автонумерации; маркированные и прочие абзацы дописываются к предыдущему элементу."""
    paras = [p for p in paras if p.text]
    texted = any(_NUMBER.match(p.text) for p in paras)
    items: list[list] = []
    for p in paras:
        m = _NUMBER.match(p.text)
        if texted and m:
            items.append([int(m[1]), p.text[m.end() :]])
        elif not texted and p.numbered and p.numbered != "bullet":
            items.append([len(items) + 1, p.text])
        elif items:
            items[-1][1] += " " + p.text
        elif not texted and not any(q.numbered for q in paras):
            items.append([len(items) + 1, p.text])  # список без номеров – по порядку
    return [(n, " ".join(t.split())) for n, t in items]


def _numbers(text: str, count: int) -> list[int]:
    """Номера из ответа: «1342», «1, 3, 4, 2», «1-3-4-2»."""
    digits = re.findall(r"\d+", text)
    if len(digits) == 1 and count < 10 and len(digits[0]) == count:
        return [int(ch) for ch in digits[0]]
    return [int(d) for d in digits]


def _matching(task: FundTask, body: list[Para | Grid], answer: str, grid: Grid | None) -> None:
    table = next((b for b in body if isinstance(b, Grid) and len(b.rows[0]) >= 2), None)
    if table is None:
        raise ValueError("нет таблицы с элементами для сопоставления")
    task.text = "\n".join(
        b.text for b in body[: body.index(table)] if isinstance(b, Para) and b.text
    )
    rows = table.rows[1:] if not _LETTER.match(table.texts(0)[0]) else table.rows
    lefts: list[tuple[str, str]] = []
    for row in rows:
        text = " ".join(p.text for p in row[0] if p.text)
        if text:
            m = _LETTER.match(text)
            letter = m[1] if m else LETTERS[len(lefts)]
            lefts.append((letter.upper(), text[m.end() :] if m else text))
    rights = dict(_items([p for row in rows for cell in row[1:2] for p in cell]))
    if grid is not None and len(grid.rows) >= 2:
        pairs = dict(zip(grid.texts(0), grid.texts(1), strict=False))
        pairs = {k.strip().upper(): v for k, v in pairs.items() if k.strip()}
    else:
        pairs = {m[0].upper(): m[1] for m in re.findall(r"([А-ЯЁA-Z])\s*[-–:]?\s*(\d+)", answer)}
    if len(lefts) < 2:
        raise ValueError("в соответствии меньше двух элементов левого столбца")
    if not all(t for _, t in lefts) or not all(rights.values()):
        raise ValueError("пустой элемент в таблице соответствия")
    numbers = sorted(rights)
    options = [schemas.AnswerOptionIn(text=t) for _, t in lefts]
    options += [schemas.AnswerOptionIn(text=rights[n]) for n in numbers]
    for i, (letter, _) in enumerate(lefts):
        digits = re.findall(r"\d+", pairs.get(letter, ""))
        if not digits or int(digits[0]) not in rights:
            raise ValueError(f"в ответе нет пары для {letter}")
        options[i].match_index = len(lefts) + numbers.index(int(digits[0]))
    targets = [o.match_index for o in options if o.match_index is not None]
    if len(set(targets)) != len(targets):
        raise ValueError("в ответе один номер у двух букв")
    task.options = options


def _sequence(task: FundTask, body: list[Para | Grid], answer: str, grid: Grid | None) -> None:
    paras = [b for b in body if isinstance(b, Para) and b.text]
    paras = [p for p in paras if not _norm(p.text).startswith("запишите")]
    first = next(
        (
            i
            for i, p in enumerate(paras)
            if (p.numbered and p.numbered != "bullet") or _NUMBER.match(p.text)
        ),
        None,
    )
    if first is None:
        raise ValueError("нет пронумерованных элементов последовательности")
    task.text = "\n".join(p.text for p in paras[:first])
    items = dict(_items(paras[first:]))
    if len(items) < 2:
        raise ValueError("в последовательности меньше двух элементов")
    order = _numbers(" ".join(grid.texts(0)) if grid is not None else answer, len(items))
    if sorted(order) != sorted(items):
        raise ValueError("ответ – не перестановка номеров элементов")
    task.options = [schemas.AnswerOptionIn(text=items[n]) for n in order]


def _open(task: FundTask, body: list[Para | Grid], answer: str, grid: Grid | None) -> None:
    lines = []
    for b in body:
        if isinstance(b, Para) and b.text:
            lines.append(f"– {b.text}" if b.numbered == "bullet" else b.text)
    task.text = "\n".join(lines)
    if grid is not None and not answer:
        answer = " ".join(t for r in range(len(grid.rows)) for t in grid.texts(r) if t)
    answers = [a.strip().rstrip(".").strip() for a in re.split(r"\s/\s|;", answer)]
    answers = [a for a in answers if a and not _BLANK.match(a)]
    if not answers:
        raise ValueError("нет ответа")
    task.options = [schemas.AnswerOptionIn(text=a) for a in answers]


_PARSERS = {
    models.QuestionKind.MATCHING: _matching,
    models.QuestionKind.SEQUENCE: _sequence,
    models.QuestionKind.OPEN: _open,
}


def parse(filename: str, data: bytes) -> FundDocument:
    """Разбирает документ; если это не ФОС или файл поврежден – документ с error (в архиве
    остальные документы загружаются)."""
    try:
        with files.reading():
            blocks = _blocks(data)
    except DomainError as exc:
        return FundDocument(filename, "", None, None, {}, [], [], error=exc.detail)
    header, distribution = None, None
    tasks: list[FundTask] = []
    warnings: list[str] = []
    section: str | None = None
    current: FundTask | None = None
    body: list[Para | Grid] = []
    answer: str | None = None  # текст после «Ответ:», пока ждем таблицу ответа

    def finish(grid: Grid | None = None) -> None:
        nonlocal current, body, answer
        if current is not None:
            try:
                if answer is None:
                    raise ValueError("нет поля «Ответ»")
                _PARSERS[current.kind](current, body, answer, grid)
                if not current.text:
                    raise ValueError("нет текста задания")
            except (ValueError, ValidationError) as exc:
                current.error = str(exc).splitlines()[0]
            tasks.append(current)
        current, body, answer = None, [], None

    for i, block in enumerate(blocks):
        if isinstance(block, Grid):
            if not block.rows or not block.rows[0]:
                continue
            first = _norm(block.texts(0)[0])
            if header is None and current is None and (parsed := _header(block)) is not None:
                header = parsed
            elif (rows := _distribution(block)) is not None:
                finish()
                distribution = rows
            elif first.startswith("сценарий выполнения"):
                continue
            elif current is not None and answer is not None:
                finish(block)  # таблица ответа
            elif current is not None:
                body.append(block)
            continue
        text = block.text
        kind = _task_kind(text) if text else None
        if kind is not None:
            finish()
            number = sum(1 for t in tasks if t.competency == section) + 1
            current = FundTask(number=number, competency=section, kind=kind)
            continue
        m = _COMPETENCY.match(text)
        if m and len(m[3]) > 10 and (current is None or answer is not None):
            finish()
            section = _code(m[1], m[2])
            continue
        if current is None:
            continue
        m = _ANSWER.search(text) if answer is None else None
        if m:
            # «Ответ:» – отдельным абзацем или в конце абзаца с текстом задания.
            if text[: m.start()].strip():
                body.append(Para(text[: m.start()].strip(), block.numbered))
            answer = m[1].strip()
            following = next(
                (b for b in blocks[i + 1 :] if not isinstance(b, Para) or b.text), None
            )
            if answer or not isinstance(following, Grid):
                finish()
        elif answer is None:
            body.append(block)
    finish()

    if header is None:
        return FundDocument(
            filename, "", None, None, {}, [], [],
            error="нет таблицы «Информация о дисциплине» – это не ФОС дисциплины",
        )  # fmt: skip
    direction, semester, discipline, competencies = header
    _apply_distribution(tasks, distribution, warnings)
    for t in tasks:
        if t.competency is None and len(competencies) == 1:
            t.competency = next(iter(competencies))
    return FundDocument(filename, direction, semester, discipline, competencies, tasks, warnings)


def _apply_distribution(tasks: list[FundTask], rows, warnings: list[str]) -> None:
    """Уровень сложности (и компетенция, если раздела нет) – из таблицы распределения:
    задания идут в ней в том же порядке, по номерам внутри компетенции или подряд."""
    if not rows:
        if tasks:
            warnings.append("нет таблицы распределения заданий – уровень сложности «базовый»")
        return
    if len(rows) == len(tasks):
        for task, (code, _, difficulty) in zip(tasks, rows, strict=True):
            task.competency = task.competency or code
            task.difficulty = difficulty or task.difficulty
        return
    by_number = {(code, number): difficulty for code, number, difficulty in rows}
    for task in tasks:
        difficulty = by_number.get((task.competency, task.number))
        if difficulty is not None:
            task.difficulty = difficulty
    warnings.append(
        f"заданий в документе {len(tasks)}, а в таблице распределения {len(rows)} – уровень "
        "сложности взят по номерам заданий"
    )


def documents(filename: str, data: bytes) -> list[tuple[str, bytes]]:
    """Документы из загруженного файла: сам .docx или все .docx из .zip."""
    ext = files.file_kind(filename, data, (".docx", ".zip"))
    if ext == ".docx":
        return [(filename, data)]
    result = []
    with files.reading(), zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = [i for i in archive.infolist() if not i.is_dir()]
        if sum(i.file_size for i in entries) > MAX_UNPACKED:
            raise DomainError("архив слишком большой: в распакованном виде больше 50 МБ")
        for info in entries:
            # Имена из архивов Windows – в кодировке cp866, если не помечены как UTF-8.
            name = info.filename
            if not info.flag_bits & 0x800:
                try:
                    name = name.encode("cp437").decode("cp866")
                except UnicodeError:
                    pass
            name = PurePath(name).name
            if name.lower().endswith(".docx") and not name.startswith(("~$", ".")):
                result.append((name, archive.read(info)))
    if not result:
        raise DomainError("в архиве нет файлов .docx")
    if len(result) > MAX_FILES:
        raise DomainError(f"в архиве больше {MAX_FILES} документов – загрузите по частям")
    return sorted(result)


# --- сопоставление с базой и сохранение -------------------------------------------------


@dataclass
class _Target:
    """Куда загружается документ и что там уже есть."""

    direction: models.Direction | None
    discipline: models.Discipline | None
    competencies: dict[str, models.Competency]
    existing: set[tuple[int, tuple]]  # (id компетенции, отпечаток вопроса) – уже в банке
    error: str | None = None
    warning: str | None = None


def _plan_of_document(
    session: Session, doc: FundDocument, direction_id: int | None
) -> tuple[models.Direction | None, str | None, str | None]:
    """План, куда идут вопросы: (план, ошибка, предупреждение). Выбранный в шапке кабинета
    важнее, пока код документа ему не противоречит."""
    selected = get_or_404(session, models.Direction, direction_id) if direction_id else None
    if not doc.direction_code or (selected and selected.code == doc.direction_code):
        if selected:
            return selected, None, None
        return (
            None,
            "учебный план (в шапке нет кода) не найден – выберите его в шапке кабинета",
            None,
        )
    plans = list(
        session.scalars(select(models.Direction).where(models.Direction.code == doc.direction_code))
    )
    if not plans:
        return (
            None,
            f"учебный план {doc.direction_code} не найден – загрузите его или добавьте "
            "в «Учебных планах»",
            None,
        )
    if len(plans) > 1:
        return (
            None,
            f"у кода {doc.direction_code} несколько учебных планов – выберите план в шапке",
            None,
        )
    warning = None
    if selected:
        warning = (
            f"код документа не совпадает с выбранным планом «{selected.label}»: "
            f"вопросы пойдут в «{plans[0].label}»"
        )
    return plans[0], None, warning


def _target(
    session: Session,
    doc: FundDocument,
    user: models.User,
    direction_id: int | None,
    discipline_id: int | None,
) -> _Target:
    if doc.error:
        return _Target(None, None, {}, set(), doc.error)
    warning = None
    if discipline_id is not None:
        discipline = get_or_404(session, models.Discipline, discipline_id)
        if user.role != models.Role.ADMIN:
            teaching.own_discipline(session, user, discipline.id)
        direction = discipline.direction
    else:
        direction, error, warning = _plan_of_document(session, doc, direction_id)
        if direction is None:
            return _Target(None, None, {}, set(), error)
        if not doc.discipline:
            return _Target(direction, None, {}, set(), "в шапке документа не указана дисциплина")
        discipline = next(
            (
                d
                for d in session.scalars(
                    select(models.Discipline).where(models.Discipline.direction_id == direction.id)
                )
                if _norm(d.name) == _norm(doc.discipline)
            ),
            None,
        )
    competencies = {
        c.code: c
        for c in session.scalars(
            select(models.Competency).where(models.Competency.direction_id == direction.id)
        )
    }
    existing: set[tuple[int, tuple]] = set()
    if discipline is not None and discipline.id is not None:
        existing = {
            (q.competency_id, _question_print(q))
            for q in session.scalars(
                select(models.Question)
                .where(models.Question.discipline_id == discipline.id)
                .options(selectinload(models.Question.options))
            )
        }
    return _Target(direction, discipline, competencies, existing, warning=warning)


# Отпечаток вопроса: текст, тип и ответ. Одинаковый общий текст («Соотнесите понятия с
# определениями») с другими парами – другой вопрос, а не дубль.


def _print(text: str, kind, pairs, items) -> tuple:
    pairs = tuple(sorted((_norm(a), _norm(b)) for a, b in pairs))
    items = [_norm(i) for i in items]
    ordered = tuple(items) if kind == models.QuestionKind.SEQUENCE else tuple(sorted(items))
    return _norm(text), kind, pairs, ordered


def _task_print(task: FundTask) -> tuple:
    options = task.options
    pairs = [(o.text, options[o.match_index].text) for o in options if o.match_index is not None]
    items = [o.text for o in options if o.match_index is None]
    return _print(task.text, task.kind, pairs, items)


def _question_print(q: models.Question) -> tuple:
    pairs = [(o.text, o.match.text) for o in q.options if o.match is not None]
    items = [o.text for o in q.options if o.match is None]
    return _print(q.text, q.kind, pairs, items)


def _answer(task: FundTask) -> list[str]:
    options = task.options
    if task.kind == models.QuestionKind.MATCHING:
        return [
            f"{o.text} → {options[o.match_index].text}"
            for o in options
            if o.match_index is not None
        ]
    return [o.text for o in options]


def _check(
    task: FundTask, target: _Target, doc: FundDocument, user: models.User, seen: set
) -> tuple[str, str | None]:
    """Статус задания: new, duplicate или error с причиной."""
    if task.error:
        return "error", task.error
    if task.competency is None:
        return "error", "не указана компетенция (нет раздела с кодом и таблицы распределения)"
    competency = target.competencies.get(task.competency)
    if competency is None and (
        user.role != models.Role.ADMIN or not doc.competencies.get(task.competency)
    ):
        return "error", f"{task.competency} нет в направлении"
    if user.role != models.Role.ADMIN and target.discipline is not None:
        if competency.id not in target.discipline.competency_ids:
            return "error", f"дисциплина не формирует {task.competency}"
    key = (task.competency, _task_print(task))
    if key in seen or (competency is not None and (competency.id, key[1]) in target.existing):
        return "duplicate", None
    seen.add(key)
    return "new", None


def preview(
    session: Session,
    docs: list[FundDocument],
    user: models.User,
    direction_id: int | None = None,
    discipline_id: int | None = None,
) -> schemas.FundPreview:
    result = []
    for doc in docs:
        target = _target(session, doc, user, direction_id, discipline_id)
        error = target.error
        if error is None and target.discipline is None and user.role != models.Role.ADMIN:
            error = f"дисциплины «{doc.discipline}» нет – ее добавляет администратор"
        seen: set = set()
        questions = []
        for task in doc.tasks:
            status, reason = ("error", error) if error else _check(task, target, doc, user, seen)
            questions.append(
                schemas.FundQuestionItem(
                    number=task.number,
                    competency=task.competency,
                    kind=task.kind,
                    difficulty=task.difficulty,
                    text=task.text,
                    answer=_answer(task) if not task.error else [],
                    status=status,
                    error=reason,
                )
            )
        result.append(
            schemas.FundDocumentPreview(
                filename=doc.filename,
                direction=doc.direction,
                direction_id=target.direction.id if target.direction else None,
                target_label=target.direction.label if target.direction else None,
                discipline=target.discipline.name if target.discipline else doc.discipline,
                discipline_exists=target.discipline is not None,
                semester=doc.semester,
                competencies=[
                    schemas.FundCompetencyItem(
                        code=code, description=description, exists=code in target.competencies
                    )
                    for code, description in sorted(
                        doc.competencies.items(), key=lambda kv: code_key(kv[0])
                    )
                ],
                questions=questions,
                warnings=[*doc.warnings, *filter(None, [target.warning])],
                error=error,
            )
        )
    return schemas.FundPreview(documents=result)


def apply(
    session: Session,
    docs: list[FundDocument],
    user: models.User,
    direction_id: int | None = None,
    discipline_id: int | None = None,
) -> schemas.FundResult:
    """Добавляет новые вопросы; дубли и неразобранные задания пропускает. Администратору
    недостающие дисциплина и компетенции (с формулировкой из шапки) создаются."""
    result = schemas.FundResult()
    for doc in docs:
        target = _target(session, doc, user, direction_id, discipline_id)
        if target.error or (target.discipline is None and user.role != models.Role.ADMIN):
            result.errors += len(doc.tasks)
            continue
        seen: set = set()
        for task in doc.tasks:
            status, _ = _check(task, target, doc, user, seen)
            if status != "new":
                result.duplicates += status == "duplicate"
                result.errors += status == "error"
                continue
            if target.discipline is None:
                target.discipline = models.Discipline(
                    direction=target.direction, name=doc.discipline
                )
                session.add(target.discipline)
                result.created_disciplines += 1
            competency = target.competencies.get(task.competency)
            if competency is None:
                competency = models.Competency(
                    direction=target.direction,
                    code=task.competency,
                    kind=models.CompetencyKind(task.competency.split("-")[0]),
                    description=doc.competencies[task.competency],
                )
                session.add(competency)
                target.competencies[task.competency] = competency
                result.created_competencies += 1
            if competency not in target.discipline.competencies:
                target.discipline.competencies.append(competency)
            data = schemas.QuestionIn(
                competency_id=0,  # проверка состава вариантов; связи – ниже, объектами
                text=task.text,
                kind=task.kind,
                difficulty=task.difficulty,
                options=task.options,
            )
            session.add(
                models.Question(
                    competency=competency,
                    discipline=target.discipline,
                    text=task.text,
                    kind=task.kind,
                    difficulty=task.difficulty,
                    options=questions.build_options(data),
                )
            )
            result.created_questions += 1
            if target.direction.label not in result.directions:
                result.directions.append(target.direction.label)
    session.commit()
    return result
