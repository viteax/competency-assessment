"""Импорт компетенций и индикаторов из файлов ОПОП / РПД.

Структура таких документов у вузов разная (таблица «код — формулировка — индикаторы»,
сплошной текст, матрица компетенций в Excel), поэтому разбор не завязан на вёрстку:
из файла извлекается текст, а в нём ищутся коды вида УК-1, ОПК-5, ПК-2 и коды
индикаторов (УК-1.1, ИУК-1.1, ИД-1.УК-1, ИД-1 ОПК-5 …). Формулировка — текст после
кода до следующего кода. Результат показывается администратору для проверки и только
потом сохраняется.
"""

import io
import re
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app import models, schemas
from app.errors import DomainError

MAX_FILE_SIZE = 10 * 1024 * 1024
SUPPORTED = (".docx", ".xlsx", ".pdf", ".txt")

_DASH = r"[-‑–—]"
_KIND = r"(?P<{name}>УК|ОПК|ПК)"
_NOT_LETTER = r"(?<![А-Яа-яЁёA-Za-z])"

_CODE_RE = re.compile(
    "|".join(
        [
            # ИД-1.УК-1, ИД-1 ОПК-5, ИД1.ПК-2 — номер индикатора перед кодом компетенции
            rf"{_NOT_LETTER}ИДК?\s?{_DASH}?\s?(?P<a_ind>\d+)\s?[.\s]?\s?"
            rf"{_KIND.format(name='a_kind')}\s?{_DASH}\s?(?P<a_num>\d+)",
            # УК-1.1, ИУК-1.1, ИД-УК-1.1, ИДК-ОПК-5.2
            rf"{_NOT_LETTER}(?:ИДК?\s?{_DASH}?\s?|И)?{_KIND.format(name='b_kind')}"
            rf"\s?{_DASH}\s?(?P<b_num>\d+)\.(?P<b_ind>\d+)",
            # УК-1, ОПК-5
            rf"{_NOT_LETTER}{_KIND.format(name='c_kind')}\s?{_DASH}\s?(?P<c_num>\d+)(?![.,]?\d)",
        ]
    )
)

_DISCIPLINE_RE = re.compile(
    r"дисциплин[аыу]?\s*(?:\(\s*модул[яь]\s*\)\s*)?[«\"“](?P<name>[^»\"”\n]{2,200})[»\"”]",
    re.IGNORECASE,
)

# Подписи колонок и служебные слова, которые не являются формулировкой.
_NOISE_RE = re.compile(
    r"^(код и наименование.*|наименование.*компетенци.*|индикатор.*достижени.*компетенци.*)$",
    re.IGNORECASE,
)


@dataclass
class ParsedCompetency:
    code: str
    kind: models.CompetencyKind
    description: str = ""
    indicators: dict[str, str] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    discipline_name: str | None
    competencies: dict[str, ParsedCompetency]


# --- извлечение текста ------------------------------------------------------


def _docx_text(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    chunks: list[str] = []

    def table_text(table: Table) -> None:
        for row in table.rows:
            previous = None
            for cell in row.cells:
                # Объединённые ячейки python-docx возвращает несколько раз подряд.
                if cell._tc is previous:
                    continue
                previous = cell._tc
                chunks.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            chunks.append(block.text)
        elif isinstance(block, Table):
            table_text(block)
    return "\n".join(chunks)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            chunks.extend(str(v) for v in row if v not in (None, ""))
    return "\n".join(chunks)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if len(data) > MAX_FILE_SIZE:
        raise DomainError("файл больше 10 МБ")
    if name.endswith((".doc", ".xls")):
        raise DomainError("старые форматы .doc/.xls не поддерживаются — сохраните как .docx/.xlsx")
    try:
        if name.endswith(".docx"):
            return _docx_text(data)
        if name.endswith(".xlsx"):
            return _xlsx_text(data)
        if name.endswith(".pdf"):
            return _pdf_text(data)
        if name.endswith(".txt"):
            return data.decode("utf-8-sig", errors="replace")
    except DomainError:
        raise
    except Exception as exc:  # повреждённый или не тот файл
        raise DomainError(f"не удалось прочитать файл: {exc}") from exc
    raise DomainError(f"поддерживаются файлы {', '.join(SUPPORTED)}")


# --- разбор текста ----------------------------------------------------------


def _clean(text: str) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    text = text.strip(" .:;,–—-|")
    text = text.removeprefix("– ").removeprefix("— ")
    if _NOISE_RE.match(text):
        return ""
    return text[:2000]


def _better(current: str, candidate: str) -> str:
    """Из нескольких упоминаний кода оставляем самую содержательную формулировку."""
    if not candidate:
        return current
    if not current:
        return candidate
    return candidate if len(candidate) > len(current) else current


def parse_text(text: str) -> ParsedDocument:
    matches = list(_CODE_RE.finditer(text))
    competencies: dict[str, ParsedCompetency] = {}

    def competency(kind: str, num: str) -> ParsedCompetency:
        code = f"{kind}-{int(num)}"
        if code not in competencies:
            competencies[code] = ParsedCompetency(code=code, kind=models.CompetencyKind(kind))
        return competencies[code]

    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        # Формулировка не перетекает через пустую строку или в следующий абзац списка кодов.
        tail = text[m.end() : end]
        description = _clean(re.split(r"\n\s*\n", tail, maxsplit=1)[0])
        g = m.groupdict()
        if g["a_kind"]:
            parent = competency(g["a_kind"], g["a_num"])
            ind_code = f"{parent.code}.{int(g['a_ind'])}"
        elif g["b_kind"]:
            parent = competency(g["b_kind"], g["b_num"])
            ind_code = f"{parent.code}.{int(g['b_ind'])}"
        else:
            parent = competency(g["c_kind"], g["c_num"])
            parent.description = _better(parent.description, description)
            continue
        parent.indicators[ind_code] = _better(parent.indicators.get(ind_code, ""), description)

    discipline = _DISCIPLINE_RE.search(text)
    return ParsedDocument(
        discipline_name=_clean(discipline.group("name")) if discipline else None,
        competencies=dict(sorted(competencies.items(), key=lambda kv: _sort_key(kv[0]))),
    )


def _sort_key(code: str) -> tuple[int, int]:
    kind, num = code.split("-")
    return (["УК", "ОПК", "ПК"].index(kind), int(num))


# --- предпросмотр и сохранение ---------------------------------------------


def _existing(session: Session, codes: list[str]) -> dict[str, models.Competency]:
    rows = session.scalars(
        select(models.Competency)
        .where(models.Competency.code.in_(codes))
        .options(selectinload(models.Competency.indicators))
    )
    return {c.code: c for c in rows}


def preview(session: Session, filename: str, parsed: ParsedDocument) -> schemas.ImportPreview:
    existing = _existing(session, list(parsed.competencies))
    items = []
    warnings = []
    for code, pc in parsed.competencies.items():
        current = existing.get(code)
        current_indicators = {i.code: i.description for i in current.indicators} if current else {}
        indicators = [
            schemas.ImportIndicator(
                code=ic,
                description=desc,
                exists=ic in current_indicators,
                changed=ic in current_indicators and bool(desc) and desc != current_indicators[ic],
            )
            for ic, desc in sorted(pc.indicators.items(), key=lambda kv: int(kv[0].rsplit(".")[1]))
        ]
        if current is None:
            status = "new"
        elif any(not i.exists or i.changed for i in indicators) or (
            pc.description and pc.description != current.description
        ):
            status = "update"
        else:
            status = "same"
        if current is None and not pc.description:
            warnings.append(f"{code}: формулировка не найдена — укажите её перед импортом")
        items.append(
            schemas.ImportCompetency(
                code=code,
                kind=pc.kind,
                description=pc.description,
                existing_description=current.description if current else None,
                status=status,
                indicators=indicators,
            )
        )
    discipline_exists = bool(
        parsed.discipline_name
        and session.scalar(
            select(models.Discipline.id).where(models.Discipline.name == parsed.discipline_name)
        )
    )
    if not items:
        warnings.append("в файле не найдено кодов компетенций (УК-…, ОПК-…, ПК-…)")
    return schemas.ImportPreview(
        filename=filename,
        discipline_name=parsed.discipline_name,
        discipline_exists=discipline_exists,
        competencies=items,
        warnings=warnings,
    )


def apply(session: Session, data: schemas.ImportApply) -> schemas.ImportResult:
    """Создаёт недостающие компетенции и индикаторы, при желании обновляет формулировки,
    и привязывает компетенции к дисциплине. Ничего не удаляет."""
    existing = _existing(session, [c.code for c in data.competencies])
    result = schemas.ImportResult()
    touched: list[models.Competency] = []

    for item in data.competencies:
        competency = existing.get(item.code)
        if competency is None:
            if not item.description:
                raise DomainError(f"{item.code}: укажите формулировку компетенции")
            competency = models.Competency(
                code=item.code,
                kind=models.CompetencyKind(item.code.split("-")[0]),
                description=item.description,
            )
            session.add(competency)
            result.created_competencies += 1
        elif (
            data.overwrite_descriptions
            and item.description
            and item.description != competency.description
        ):
            competency.description = item.description
            result.updated_competencies += 1
        touched.append(competency)

        indicators = {i.code: i for i in competency.indicators}
        for ind in item.indicators:
            current = indicators.get(ind.code)
            if current is None:
                competency.indicators.append(
                    models.Indicator(code=ind.code, description=ind.description or ind.code)
                )
                result.created_indicators += 1
            elif (
                data.overwrite_descriptions
                and ind.description
                and ind.description != current.description
            ):
                current.description = ind.description
                result.updated_indicators += 1

    discipline = None
    if data.discipline_name:
        discipline = session.scalar(
            select(models.Discipline).where(models.Discipline.name == data.discipline_name)
        )
        if discipline is None:
            discipline = models.Discipline(name=data.discipline_name)
            session.add(discipline)
            result.created_discipline = True
        linked = {c.code for c in discipline.competencies}
        for competency in touched:
            if competency.code not in linked:
                discipline.competencies.append(competency)
                result.linked_competencies += 1

    session.commit()
    if discipline is not None:
        result.discipline = schemas.DisciplineOut.model_validate(discipline)
    return result
