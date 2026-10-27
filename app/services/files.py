"""Чтение загруженных файлов: проверка размера и формата, текст и таблицы из
.docx / .xlsx / .pdf / .csv / .txt. Общее для импорта компетенций и студентов."""

import csv
import io
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import PurePath

from app.errors import DomainError

MAX_FILE_SIZE = 10 * 1024 * 1024
_OLD_FORMATS = {".doc": ".docx", ".xls": ".xlsx"}


def file_kind(filename: str, data: bytes, allowed: tuple[str, ...]) -> str:
    """Проверяет файл и возвращает его расширение."""
    ext = PurePath(filename.lower()).suffix
    if len(data) > MAX_FILE_SIZE:
        raise DomainError("файл больше 10 МБ")
    if ext in _OLD_FORMATS:
        raise DomainError(
            f"старый формат {ext} не поддерживается – сохраните файл как {_OLD_FORMATS[ext]}"
        )
    if ext not in allowed:
        raise DomainError(f"поддерживаются файлы {', '.join(allowed)}")
    return ext


@contextmanager
def reading() -> Iterator[None]:
    """Повреждённый или «не тот» файл превращается в понятную ошибку 400."""
    try:
        yield
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(f"не удалось прочитать файл: {exc}") from exc


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DomainError("не удалось определить кодировку – сохраните файл в UTF-8")


def csv_rows(text: str) -> list[list[str]]:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 2.0 → «2»: числа из Excel – как в таблице
    return str(value)


def xlsx_sheets(data: bytes) -> dict[str, list[list[str]]]:
    """Все листы книги: название → строки, ячейки – текстом (пустые – "")."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return {
        sheet.title: [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        for sheet in workbook.worksheets
    }


def xlsx_rows(data: bytes) -> list[list[str]]:
    """Строки первого листа."""
    sheets = xlsx_sheets(data)
    return next(iter(sheets.values()), [])


def docx_text(data: bytes) -> str:
    """Текст абзацев и ячеек таблиц в порядке документа."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    chunks: list[str] = []

    def table_text(table: Table) -> None:
        for row in table.rows:
            previous = None
            for cell in row.cells:
                # Объединённые ячейки python-docx возвращает несколько раз подряд.
                if cell._tc is previous:
                    continue
                previous = cell._tc
                chunks.append(cell.text)
                for nested in cell.tables:
                    table_text(nested)

    for block in Document(io.BytesIO(data)).iter_inner_content():
        if isinstance(block, Paragraph):
            chunks.append(block.text)
        elif isinstance(block, Table):
            table_text(block)
    return "\n".join(chunks)


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
