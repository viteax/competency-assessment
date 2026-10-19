"""Выгрузка результатов группы в Excel: компетенции и трудовые функции по каждому студенту."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app import models
from app.services import app_settings, reports

GREEN = PatternFill("solid", fgColor="E5F5EC")
RED = PatternFill("solid", fgColor="FDEEE5")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws: Worksheet, titles: list[str], widths: list[int]) -> None:
    ws.append(titles)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(row=1, column=i).font = BOLD
        ws.cell(row=1, column=i).alignment = WRAP
    ws.freeze_panes = "B2"


def group_results_xlsx(session: Session, group_id: int, test_id: int | None = None) -> bytes:
    report = reports.group_report(session, group_id, test_id)
    threshold = app_settings.mastery_threshold(session)
    test = session.get(models.Test, test_id) if test_id is not None else None
    competencies = sorted(
        {(c.code, c.competency_id) for s in report.students for c in s.competencies}
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Компетенции"
    _header(
        ws,
        ["ФИО", "Группа", "Дата тестирования", *[code for code, _ in competencies]],
        [36, 12, 18, *[11] * len(competencies)],
    )
    for s in report.students:
        by_id = {c.competency_id: c for c in s.competencies}
        ws.append(
            [
                s.full_name,
                report.name,
                s.tested_at.astimezone().strftime("%d.%m.%Y") if s.tested_at else "",
                *[
                    f"{by_id[cid].correct} из {by_id[cid].total}" if cid in by_id else ""
                    for _, cid in competencies
                ],
            ]
        )
        for col, (_, cid) in enumerate(competencies, start=4):
            if cid in by_id:
                ws.cell(row=ws.max_row, column=col).fill = GREEN if by_id[cid].mastered else RED
    ws.append([])
    ws.append(
        [
            f"Тест: «{test.title}»" if test else "Все пройденные тесты",
            None,
            None,
            f"«N из M» – верных ответов из заданных. Зелёным – компетенция сформирована "
            f"(от {round(threshold * 100)}% верных), красным – нет.",
        ]
    )

    # Уровень ТФ студента – средняя доля верных ответов по её проверенным компетенциям.
    functions = reports.labor_functions(session)
    tf = wb.create_sheet("Трудовые функции")
    _header(
        tf,
        ["ФИО", "Группа", *[f"{lf.code} {lf.name}" for lf in functions]],
        [36, 12, *[18] * len(functions)],
    )
    tf.row_dimensions[1].height = 75
    for s in report.students:
        level = {c.competency_id: c.ratio for c in s.competencies}
        axes = reports.labor_function_axes(functions, level)
        tf.append(
            [
                s.full_name,
                report.name,
                *[f"{round(a.value * 100)}%" if a.value is not None else "" for a in axes],
            ]
        )
    tf.append([])
    tf.append(
        [
            "Уровень ТФ – средняя доля верных ответов по связанным с ней компетенциям, "
            "которые проверялись; пусто – ни одна не проверялась."
        ]
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
