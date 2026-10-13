"""Импорт списка студентов из выгрузки СДО (CSV / Excel).

Формат выгрузки заранее неизвестен (участники курса Moodle, журнал оценок, список из
деканата), поэтому колонки определяются по заголовкам, а администратор может переназначить
их в предпросмотре. Студенты сопоставляются с уже существующими по идентификатору СДО,
затем по e-mail, затем по ФИО в той же группе.
"""

import csv
import io
import re
import secrets
from dataclasses import dataclass

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models, schemas
from app.errors import DomainError
from app.security import hash_password
from app.services.common import get_or_404
from app.services.importer import MAX_FILE_SIZE

FIELDS = (
    "full_name",
    "last_name",
    "first_name",
    "middle_name",
    "email",
    "external_id",
    "login",
    "group",
    "role",
)

# Порядок важен: более конкретные заголовки проверяются раньше общих.
_HEADER_RULES: list[tuple[str, re.Pattern[str]]] = [
    (
        "full_name",
        re.compile(
            r"^(фио|ф\.\s?и\.\s?о\.?|имя\s*/\s*фамилия|фамилия\s*/?\s*имя|"
            r"полное имя|full ?name|студент|обучающийся|слушатель)"
        ),
    ),
    ("email", re.compile(r"(почт|e-?mail)")),
    (
        "external_id",
        re.compile(
            r"(индивидуальный номер|id ?number|idnumber|зач[её]тн|"
            r"номер студ|табельный|^id$)"
        ),
    ),
    ("login", re.compile(r"(логин|username|имя пользователя)")),
    ("middle_name", re.compile(r"(отчество|middle ?name)")),
    ("last_name", re.compile(r"(фамилия|last ?name|surname)")),
    ("first_name", re.compile(r"(^имя|first ?name)")),
    ("group", re.compile(r"(групп|group|^отдел|department)")),
    ("role", re.compile(r"(^рол|^role)")),
]

_STUDENT_ROLE = re.compile(r"(студент|student|обучающ|слушател)", re.IGNORECASE)


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DomainError("не удалось определить кодировку CSV — сохраните файл в UTF-8")


def read_table(filename: str, data: bytes) -> Table:
    name = filename.lower()
    if len(data) > MAX_FILE_SIZE:
        raise DomainError("файл больше 10 МБ")
    try:
        if name.endswith((".csv", ".txt")):
            text = _decode(data)
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
            raw = list(csv.reader(io.StringIO(text), dialect))
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook

            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            raw = [
                ["" if v is None else str(v) for v in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        elif name.endswith(".xls"):
            raise DomainError("старый формат .xls не поддерживается — сохраните как .xlsx или CSV")
        else:
            raise DomainError("поддерживаются файлы .csv и .xlsx")
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(f"не удалось прочитать файл: {exc}") from exc

    rows = [[c.strip() for c in r] for r in raw if any(c.strip() for c in r)]
    if not rows:
        raise DomainError("файл пустой")
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return Table(headers=rows[0], rows=rows[1:])


def guess_mapping(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = dict.fromkeys(FIELDS)
    for index, header in enumerate(headers):
        h = re.sub(r"\s+", " ", header.lower()).strip()
        for field, rule in _HEADER_RULES:
            if mapping[field] is None and rule.search(h):
                mapping[field] = index
                break
    # Если ФИО собрано в одну колонку, отдельные части не нужны.
    if mapping["full_name"] is not None:
        mapping["last_name"] = mapping["first_name"] = mapping["middle_name"] = None
    return mapping


def _cell(row: list[str], index: int | None) -> str:
    return row[index].strip() if index is not None and index < len(row) else ""


def _parse_row(row: list[str], m: dict[str, int | None]) -> dict[str, str]:
    full_name = _cell(row, m["full_name"]) or " ".join(
        p
        for p in (
            _cell(row, m["last_name"]),
            _cell(row, m["first_name"]),
            _cell(row, m["middle_name"]),
        )
        if p
    )
    group = _cell(row, m["group"])
    # В Moodle в колонке «Группы» может быть несколько групп через запятую — берём первую.
    group = re.split(r"[,;]", group)[0].strip()
    return {
        "full_name": re.sub(r"\s+", " ", full_name).strip(),
        "email": _cell(row, m["email"]).lower(),
        "external_id": _cell(row, m["external_id"]) or _cell(row, m["login"]),
        "login": _cell(row, m["login"]),
        "group": group,
        "role": _cell(row, m["role"]),
    }


def name_key(full_name: str) -> tuple[str, ...]:
    """Ключ для сравнения ФИО без учёта порядка слов, регистра и ё/е.

    Moodle часто выводит «Имя Отчество Фамилия», а в деканатских списках «Фамилия Имя
    Отчество» — это один и тот же человек.
    """
    return tuple(sorted(full_name.lower().replace("ё", "е").split()))


def _find_existing(session: Session, row: schemas.StudentImportRowIn) -> models.Student | None:
    if row.external_id:
        found = session.scalar(
            select(models.Student).where(models.Student.external_id == row.external_id)
        )
        if found:
            return found
    if row.email:
        found = session.scalar(select(models.Student).where(models.Student.email == row.email))
        if found:
            return found
    if row.group:
        key = name_key(row.full_name)
        candidates = session.scalars(
            select(models.Student).join(models.Group).where(models.Group.name == row.group)
        )
        return next((s for s in candidates if name_key(s.full_name) == key), None)
    return None


def preview(
    session: Session, filename: str, table: Table, mapping: dict[str, int | None] | None
) -> schemas.StudentImportPreview:
    m = guess_mapping(table.headers) if mapping is None else {**dict.fromkeys(FIELDS), **mapping}
    groups = {g.name: g for g in session.scalars(select(models.Group))}
    rows = []
    seen_ids: set[str] = set()
    for number, raw in enumerate(table.rows, start=2):
        parsed = _parse_row(raw, m)
        problems = []
        if not parsed["full_name"]:
            problems.append("нет ФИО")
        if parsed["external_id"] and parsed["external_id"] in seen_ids:
            problems.append("идентификатор СДО повторяется в файле")
        seen_ids.add(parsed["external_id"])
        is_student = not parsed["role"] or bool(_STUDENT_ROLE.search(parsed["role"]))
        # Без валидации: строка без ФИО попадает в предпросмотр с пометкой о проблеме.
        item = schemas.StudentImportRowIn.model_construct(
            full_name=parsed["full_name"],
            email=parsed["email"] or None,
            external_id=parsed["external_id"] or None,
            login=parsed["login"] or None,
            group=parsed["group"] or None,
        )
        existing = _find_existing(session, item) if parsed["full_name"] else None
        if existing is None:
            status = "new"
        elif (
            (item.group and existing.group.name != item.group)
            or (item.email and existing.email != item.email)
            or (item.external_id and existing.external_id != item.external_id)
            or name_key(existing.full_name) != name_key(item.full_name)
        ):
            status = "update"
        else:
            status = "same"
        rows.append(
            schemas.StudentImportRow(
                **item.model_dump(),
                row=number,
                role=parsed["role"] or None,
                is_student=is_student,
                status=status,
                existing_student_id=existing.id if existing else None,
                existing_group=existing.group.name if existing else None,
                problems=problems,
            )
        )
    missing = sorted({r.group for r in rows if r.group and r.group not in groups})
    return schemas.StudentImportPreview(
        filename=filename,
        headers=table.headers,
        mapping=m,
        rows=rows,
        missing_groups=missing,
    )


# Без похожих символов (0/O, 1/l/I), чтобы пароль было легко переписать с листа.
_PASSWORD_ALPHABET = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def generate_password(length: int = 10) -> str:
    return "".join(secrets.choice(_PASSWORD_ALPHABET) for _ in range(length))


def _unique_username(session: Session, base: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9._-]", "", base.lower())[:48] or "student"
    candidate, n = base, 1
    while candidate in taken or session.scalar(
        select(models.User.id).where(models.User.username == candidate)
    ):
        n += 1
        candidate = f"{base}{n}"
    taken.add(candidate)
    return candidate


_TRANSLIT = str.maketrans(
    {
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "е": "e",
        "ё": "e",
        "ж": "zh",
        "з": "z",
        "и": "i",
        "й": "y",
        "к": "k",
        "л": "l",
        "м": "m",
        "н": "n",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "у": "u",
        "ф": "f",
        "х": "kh",
        "ц": "ts",
        "ч": "ch",
        "ш": "sh",
        "щ": "shch",
        "ъ": "",
        "ы": "y",
        "ь": "",
        "э": "e",
        "ю": "yu",
        "я": "ya",
    }
)


def _username_base(row: schemas.StudentImportRowIn) -> str:
    if row.login:
        return row.login
    if row.email:
        return row.email.split("@")[0]
    parts = row.full_name.lower().split()
    surname = parts[0].translate(_TRANSLIT) if parts else "student"
    initials = "".join(p[0].translate(_TRANSLIT) for p in parts[1:3])
    return f"{surname}.{initials}" if initials else surname


def apply(session: Session, data: schemas.StudentImportApply) -> schemas.StudentImportResult:
    """Создаёт и обновляет студентов, недостающие группы и (по желанию) учётные записи.
    Ничего не удаляет: студенты, которых нет в файле, остаются как есть."""
    default_group = (
        get_or_404(session, models.Group, data.default_group_id) if data.default_group_id else None
    )
    direction = (
        get_or_404(session, models.Direction, data.direction_id) if data.direction_id else None
    )
    groups = {g.name: g for g in session.scalars(select(models.Group))}
    result = schemas.StudentImportResult()
    taken_usernames: set[str] = set()
    seen_students: set[int] = set()

    for row in data.students:
        if row.group:
            group = groups.get(row.group)
            if group is None:
                if direction is None:
                    raise DomainError(
                        f"группы «{row.group}» нет — выберите направление для новых групп"
                    )
                group = models.Group(name=row.group, direction=direction)
                session.add(group)
                groups[row.group] = group
                result.created_groups.append(row.group)
        elif default_group is not None:
            group = default_group
        else:
            raise DomainError(f"{row.full_name}: не указана группа — выберите группу по умолчанию")

        student = _find_existing(session, row)
        if student is not None and student.id in seen_students:
            raise DomainError(f"{row.full_name}: студент встречается в файле дважды")
        if student is None:
            student = models.Student(
                full_name=row.full_name, group=group, email=row.email, external_id=row.external_id
            )
            session.add(student)
            result.created += 1
        else:
            changed = False
            # ФИО меняем, только если изменилось само имя, а не порядок слов.
            if name_key(student.full_name) != name_key(row.full_name):
                student.full_name = row.full_name
                changed = True
            for attr, value in (
                ("group", group),
                ("email", row.email),
                ("external_id", row.external_id),
            ):
                if value and getattr(student, attr) != value:
                    setattr(student, attr, value)
                    changed = True
            result.updated += changed
            result.unchanged += not changed
        session.flush()
        seen_students.add(student.id)

        if data.create_accounts:
            has_account = session.scalar(
                select(models.User.id).where(models.User.student_id == student.id)
            )
            if not has_account:
                username = _unique_username(session, _username_base(row), taken_usernames)
                password = generate_password()
                session.add(
                    models.User(
                        username=username,
                        password_hash=hash_password(password),
                        role=models.Role.STUDENT,
                        student_id=student.id,
                    )
                )
                result.accounts.append(
                    schemas.CreatedAccount(
                        full_name=student.full_name,
                        group=group.name,
                        username=username,
                        password=password,
                    )
                )

    session.commit()
    return result
