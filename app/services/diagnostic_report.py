"""Отчет «Результаты проверки уровня сформированности компетенций» в Word – по форме кафедры.

Шапка: группа, курс, основная профессиональная образовательная программа, дата проведения,
компетенции и трудовые функции теста. Таблица – по каждому студенту, сдавшему тест:
задания закрытого (соответствие, последовательность) и открытого типа – сколько и сколько
верно, номера неверных открытых заданий, итог и процент верных. Внизу – доля студентов,
выполнивших 70% и более заданий диагностической работы.

Номера заданий – как в документе варианта (выгрузка теста в Word): по компетенциям, внутри –
по порядку в варианте.
"""

import io
import math
from dataclasses import dataclass, field

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app import models
from app.errors import NotFound
from app.semesters import current_semester
from app.services import fund, reports
from app.services.common import code_key, get_or_404

PASS_SHARE = 0.7  # «выполнивших 70% и более заданий диагностической работы»
CLOSED = {models.QuestionKind.MATCHING, models.QuestionKind.SEQUENCE}

# Колонки таблицы формы (ширина в твипах – как в образце кафедры).
WIDTHS = [671, 1926, 1134, 1559, 1007, 1545, 2976, 907, 1361, 1574]


@dataclass
class StudentRow:
    identifier: str
    closed_total: int = 0
    closed_correct: int = 0
    open_total: int = 0
    open_correct: int = 0
    open_wrong: list[int] = field(default_factory=list)  # номера неверных открытых заданий

    @property
    def total(self) -> int:
        return self.closed_total + self.open_total

    @property
    def correct(self) -> int:
        return self.closed_correct + self.open_correct

    @property
    def percent(self) -> int:
        return math.floor(100 * self.correct / self.total + 0.5) if self.total else 0


@dataclass
class DiagnosticReport:
    group: str
    course: int | None
    program: str
    dates: str
    competencies: list[str]  # «ОПК-1 Способен …»
    labor_functions: list[str]  # «06.001 · A/01.3 Формализация …»
    rows: list[StudentRow]

    @property
    def passed_share(self) -> float:
        """Доля (в процентах) студентов, выполнивших не меньше 70% заданий."""
        if not self.rows:
            return 0.0
        passed = sum(r.total and r.correct / r.total >= PASS_SHARE for r in self.rows)
        return 100 * passed / len(self.rows)


def _identifier(student: models.Student, user: models.User | None) -> str:
    """Идентификатор студента: номер или логин в СДО, иначе логин в системе, иначе ФИО."""
    return student.external_id or (user.username if user else "") or student.full_name


def collect(session: Session, group_id: int, test_id: int) -> DiagnosticReport:
    group = get_or_404(session, models.Group, group_id)
    test = get_or_404(session, models.Test, test_id)
    assignment = session.scalar(
        select(models.Assignment).where(
            models.Assignment.group_id == group.id, models.Assignment.test_id == test.id
        )
    )
    if assignment is None:
        raise NotFound(f"тест «{test.title}» не назначен группе {group.name}")
    attempts = session.scalars(
        select(models.Attempt)
        .where(
            models.Attempt.assignment_id == assignment.id, models.Attempt.finished_at.is_not(None)
        )
        .options(selectinload(models.Attempt.answers), selectinload(models.Attempt.student))
    ).all()
    variants: dict[int, list[models.Question]] = {}
    for tq in sorted(test.test_questions, key=lambda tq: (tq.variant, tq.position)):
        variants.setdefault(tq.variant, []).append(tq.question)
    users = {
        u.student_id: u
        for u in session.scalars(
            select(models.User).where(models.User.student_id.in_([a.student_id for a in attempts]))
        )
    }

    rows = []
    for attempt in attempts:
        correct = {a.question_id for a in attempt.answers if a.is_correct}
        row = StudentRow(_identifier(attempt.student, users.get(attempt.student_id)))
        questions = fund.in_document_order(variants.get(attempt.variant, []))
        for number, q in enumerate(questions, start=1):
            right = q.id in correct
            if q.kind in CLOSED:
                row.closed_total += 1
                row.closed_correct += right
            else:
                row.open_total += 1
                row.open_correct += right
                if not right:
                    row.open_wrong.append(number)
        rows.append(row)
    rows.sort(key=lambda r: code_key(r.identifier))

    finished = sorted(a.finished_at.astimezone().date() for a in attempts)
    dates = ""
    if finished:
        first, last = finished[0].strftime("%d.%m.%Y"), finished[-1].strftime("%d.%m.%Y")
        dates = first if first == last else f"{first} – {last}"
    course = None
    if group.start_year and finished:
        course = (current_semester(group.start_year, finished[0]) + 1) // 2

    competencies = sorted(
        {q.competency for qs in variants.values() for q in qs}, key=lambda c: code_key(c.code)
    )
    tested = {c.id for c in competencies}
    functions = [
        f
        for f in reports.labor_functions(session, {group.direction_id})
        if any(c.id in tested for c in f.competencies)
    ]
    direction = group.direction
    return DiagnosticReport(
        group=group.name,
        course=course,
        program=direction.label,
        dates=dates,
        competencies=[f"{c.code} {c.description}" for c in competencies],
        labor_functions=[
            f"{f.function.general_function.standard.code} · {f.function.code} {f.function.name}"
            for f in functions
        ],
        rows=rows,
    )


# --- документ --------------------------------------------------------------------------


def _font(run, size: int, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")


def _line(doc, label: str, value: str = "") -> None:
    """Строка шапки: «Группа: ПИ-21» (подпись – полужирным)."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    _font(paragraph.add_run(label), 14, bold=True)
    _font(paragraph.add_run(f": {value}" if value else ":"), 14)


def _item(doc, text: str) -> None:
    """Компетенция или трудовая функция под строкой шапки."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    _font(paragraph.add_run(text), 12)


def _cell(cell, text: str, size: int, center: bool = True) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    _font(paragraph.add_run(text), size)


def _borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "6"), ("space", "0"), ("color", "000000")):
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _table(doc, report: DiagnosticReport) -> None:
    table = doc.add_table(rows=2 + len(report.rows), cols=len(WIDTHS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for col, width in zip(table._tbl.tblGrid.gridCol_lst, WIDTHS, strict=True):
        col.w = Twips(width)
    for row in table.rows:
        for cell, width in zip(row.cells, WIDTHS, strict=True):
            cell.width = Twips(width)
    head = {
        0: "N",
        1: "Идентификатор студента",
        7: "Общее кол-во заданий",
        8: "Общее кол-во правильно выполненных заданий",
        9: "Процент правильно выполненных заданий",
    }
    for col, text in head.items():
        _cell(table.cell(0, col).merge(table.cell(1, col)), text, 10)
    _cell(table.cell(0, 2).merge(table.cell(0, 3)), "Задания закрытого типа", 10)
    _cell(table.cell(0, 4).merge(table.cell(0, 6)), "Задания открытого типа", 10)
    sub = ["Кол-во заданий", "Кол-во правильно выполненных заданий"]
    for col, text in zip(range(2, 7), [*sub, *sub, "Номера неправильных ответов"], strict=True):
        _cell(table.cell(1, col), text, 10)
    for i, r in enumerate(report.rows, start=1):
        values = [
            str(i),
            r.identifier,
            str(r.closed_total),
            str(r.closed_correct),
            str(r.open_total),
            str(r.open_correct),
            ",".join(map(str, r.open_wrong)),
            str(r.total),
            str(r.correct),
            str(r.percent),
        ]
        for col, value in enumerate(values):
            _cell(table.cell(i + 1, col), value, 12, center=col != 1)


def document(report: DiagnosticReport) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Twips(16838), Twips(11906)
    section.top_margin, section.bottom_margin = Twips(1701), Twips(850)
    section.left_margin = section.right_margin = Twips(1134)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run("Результаты проверки уровня сформированности компетенций"), 14)
    doc.add_paragraph()
    _line(doc, "Группа", report.group)
    _line(doc, "Курс", str(report.course) if report.course else "–")
    _line(doc, "Основная профессиональная образовательная программа", report.program)
    _line(doc, "Дата проведения", report.dates or "–")
    _line(doc, "Компетенции", "" if report.competencies else "–")
    for text in report.competencies:
        _item(doc, text)
    _line(doc, "Трудовые функции", "" if report.labor_functions else "–")
    for text in report.labor_functions:
        _item(doc, text)
    doc.add_paragraph()
    _table(doc, report)
    share = doc.add_paragraph()
    _font(
        share.add_run(
            "Доля обучающихся, выполнивших 70% и более заданий диагностической работы, "
            f"составляет {report.passed_share:.2f} %"
        ),
        14,
    )
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def group_test_document(session: Session, group_id: int, test_id: int) -> tuple[bytes, str]:
    report = collect(session, group_id, test_id)
    test = session.get(models.Test, test_id)
    name = f"Результаты {report.group} – {test.title}"
    safe = "".join(ch for ch in name if ch not in '\\/:*?"<>|').strip()
    return document(report), f"{safe}.docx"
