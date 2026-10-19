"""Выгрузка заданий в Word по образцу фонда оценочных материалов (ФОМ) дисциплины.

Документ: шапка «Информация о дисциплине», задания по компетенциям (у каждого – таблица
«Сценарий выполнения / Результат оценивания» и поле «Ответ»), в конце – таблица
«Распределение тестовых заданий по видам и уровням сложности». Бланк – с пустыми полями
ответа, ключ – тот же документ с заполненными правильными ответами.

Порядок элементов последовательности и правого столбца соответствия перемешивается,
но одинаково в бланке и ключе (зависит только от id вопроса).
"""

import io
import random
from collections.abc import Iterable
from dataclasses import dataclass

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import Table

from app import models
from app.services.testing import shuffle_away

LETTERS = "АБВГДЕЖЗИКЛМНОПРСТУФХЦЧШЭЮЯ"

TITLES = {
    models.QuestionKind.MATCHING: "Задание на установление соответствия",
    models.QuestionKind.SEQUENCE: "Задание на установление последовательности",
    models.QuestionKind.OPEN: "Задание открытого типа",
}
SCORING = (
    "Полное совпадение с верным ответом – 1 балл;\nневерный ответ или его отсутствие – 0 баллов."
)
DIFFICULTY = {
    models.Difficulty.BASIC: ("базовый", "1-3"),
    models.Difficulty.ADVANCED: ("повышенный", "3-5"),
}


@dataclass
class FundHeader:
    directions: str  # «Направления (специальности)»
    semester: str
    discipline: str
    competencies: list[models.Competency]


@dataclass
class Section:
    """Компетенция и её задания."""

    competency: models.Competency
    questions: list[models.Question]


# --- оформление ---------------------------------------------------------------


def _new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(30), Mm(13)
    section.top_margin, section.bottom_margin = Mm(20), Mm(25)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    _page_number_footer(section)
    return doc


def _page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _borders(table: Table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _table(doc: Document, rows: int, widths_mm: list[float]) -> Table:
    table = doc.add_table(rows=rows, cols=len(widths_mm))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _borders(table)
    # Ширину берут из сетки таблицы (LibreOffice) и из ячеек (Word) – задаём и там, и там.
    for col, width in zip(table._tbl.tblGrid.gridCol_lst, widths_mm, strict=True):
        col.w = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm, strict=True):
            cell.width = Mm(width)
    return table


def _write(cell, text: str, bold: bool = False) -> None:
    """Текст в ячейку; переводы строк – отдельными абзацами."""
    lines = text.split("\n")
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        paragraph.add_run(line).bold = bold


def _para(doc: Document, text: str = "", bold: bool = False, center: bool = False):
    paragraph = doc.add_paragraph()
    if text:
        paragraph.add_run(text).bold = bold
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


# --- содержимое -----------------------------------------------------------------


def _header(doc: Document, header: FundHeader) -> None:
    _para(doc, "Фонд оценочных материалов дисциплины", center=True)
    _para(doc)
    _para(doc, "Информация о дисциплине", bold=True, center=True)
    table = _table(doc, 5, [42, 125])
    rows = [
        ("Направления (специальности)", header.directions or "–"),
        ("Семестр", header.semester or "–"),
        ("Дисциплина", header.discipline or "–"),
        ("Код", ", ".join(c.code for c in header.competencies) + "."),
        (
            "Наименование компетенции",
            "\n".join(f"{c.code} {c.description}" for c in header.competencies),
        ),
    ]
    for row, (label, value) in zip(table.rows, rows, strict=True):
        bold = label == "Дисциплина"
        _write(row.cells[0], label, bold=bold)
        _write(row.cells[1], value, bold=bold)


def _scenario(question: models.Question) -> str:
    kind = question.kind
    if kind == models.QuestionKind.MATCHING:
        has_extra = sum(o.match_id is None for o in question.options) > sum(
            o.match_id is not None for o in question.options
        )
        rule = (
            "Каждый элемент, представленный в правом списке (элементы списка 2), должен быть "
            "использован единожды, при этом допускается возможность полного отсутствия "
            "использования какого-либо элемента из данного списка."
            if has_extra
            else "Каждый элемент правого списка (элемент списка 2) должен быть использован "
            "один раз.\nЗапишите ответ попарно буквы и цифры вариантов ответа."
        )
        return (
            "Прочитайте задание и установите правильное соответствие.\n"
            "Сопоставьте элементы из списка 1 с элементами списка 2, чтобы сформировать "
            f"пары элементов. {rule}"
        )
    if kind == models.QuestionKind.SEQUENCE:
        return (
            "Прочитайте задание и установите правильную последовательность элементов.\n"
            "Запишите ответ в нужной последовательности без пробелов и знаков препинания."
        )
    one_word = all(" " not in o.text.strip() for o in question.options)
    what = "одно слово" if one_word else "слово или словосочетание"
    return (
        "Прочитайте задание и запишите ответ.\n"
        f"Ответом является {what}. Ответ может быть записан любым регистром."
    )


def _answer_grid(doc: Document, heads: list[str] | None, values: list[str]) -> None:
    _para(doc, "Ответ: ")
    rows = 2 if heads else 1
    table = _table(doc, rows, [20] * len(values))
    if heads:
        for cell, head in zip(table.rows[0].cells, heads, strict=True):
            _write(cell, head)
    for cell, value in zip(table.rows[-1].cells, values, strict=True):
        _write(cell, value)


def _matching(doc: Document, question: models.Question, rng: random.Random, key: bool) -> None:
    left = [o for o in question.options if o.match_id is not None]
    right = [o for o in question.options if o.match_id is None]
    rng.shuffle(right)
    number = {o.id: i + 1 for i, o in enumerate(right)}
    table = _table(doc, len(left) + 1, [55, 112])
    _write(table.cell(0, 0), "Список 1")
    _write(table.cell(0, 1), "Список 2")
    for i, o in enumerate(left):
        _write(table.cell(i + 1, 0), f"{LETTERS[i]}. {o.text}")
    merged = table.cell(1, 1).merge(table.cell(len(left), 1))
    _write(merged, "\n".join(f"{number[o.id]}. {o.text}" for o in right))
    letters = [LETTERS[i] for i in range(len(left))]
    _answer_grid(doc, letters, [str(number[o.match_id]) if key else "" for o in left])


def _sequence(doc: Document, question: models.Question, rng: random.Random, key: bool) -> None:
    items = list(question.options)  # в правильном порядке
    shown = items[:]
    shuffle_away(shown, rng)
    number = {o.id: i + 1 for i, o in enumerate(shown)}
    for o in shown:
        _para(doc, f"{number[o.id]}. {o.text}")
    _para(doc)
    _para(doc, "Запишите соответствующую последовательность цифр слева направо:")
    _answer_grid(doc, None, [str(number[o.id]) if key else "" for o in items])


def _task(doc: Document, number: int, question: models.Question, key: bool) -> None:
    _para(doc)
    _para(doc, f"{number}. {TITLES[question.kind]}", bold=True)
    table = _table(doc, 2, [100, 67])
    _write(table.cell(0, 0), "Сценарий выполнения")
    _write(table.cell(0, 1), "Результат оценивания")
    _write(table.cell(1, 0), _scenario(question))
    _write(table.cell(1, 1), SCORING)
    _para(doc)
    _para(doc, question.text).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _para(doc)
    rng = random.Random(question.id)
    if question.kind == models.QuestionKind.MATCHING:
        _matching(doc, question, rng, key)
    elif question.kind == models.QuestionKind.SEQUENCE:
        _sequence(doc, question, rng, key)
    else:
        answer = " / ".join(o.text for o in question.options)
        _para(doc, f"Ответ: {answer}." if key else "Ответ: ________________.")


def _merge_runs(table: Table, column: int, values: list[str]) -> None:
    """Объединяет подряд идущие одинаковые значения столбца (как в образце)."""
    start = 0
    for i in range(1, len(values) + 1):
        if i == len(values) or values[i] != values[start]:
            cell = table.cell(start + 1, column)
            if i - 1 > start:
                cell = cell.merge(table.cell(i, column))
            _write(cell, values[start])
            start = i


def _distribution(doc: Document, semester: str, numbered: list[tuple[int, models.Question]]):
    _para(doc)
    _para(doc, "Распределение тестовых заданий по видам и уровням сложности", bold=True)
    heads = [
        "Семестр",
        "Код компетенции",
        "Дисциплина",
        "№ задания",
        "Виды задания",
        "Уровень сложности",
        "Время выполнения",
    ]
    table = _table(doc, len(numbered) + 1, [16, 22, 33, 15, 41, 22, 18])
    for cell, head in zip(table.rows[0].cells, heads, strict=True):
        _write(cell, head, bold=True)
    _merge_runs(table, 0, [semester or "–"] * len(numbered))
    _merge_runs(table, 1, [q.competency.code for _, q in numbered])
    _merge_runs(table, 2, [q.discipline.name if q.discipline else "–" for _, q in numbered])
    for row, (number, q) in zip(table.rows[1:], numbered, strict=True):
        level, minutes = DIFFICULTY[q.difficulty]
        _write(row.cells[3], str(number))
        _write(row.cells[4], TITLES[q.kind])
        _write(row.cells[5], level)
        _write(row.cells[6], minutes)


def _fund(doc: Document, header: FundHeader, sections: list[Section], key: bool, title: str = ""):
    _header(doc, header)
    if title:
        _para(doc)
        _para(doc, title, bold=True, center=True)
    numbered: list[tuple[int, models.Question]] = []
    for section in sections:
        _para(doc)
        _para(doc, f"{section.competency.code} {section.competency.description}")
        for question in section.questions:
            numbered.append((len(numbered) + 1, question))
            _task(doc, len(numbered), question, key)
    if key:
        _para(doc)
        _para(doc, "Ключ: в полях «Ответ» указаны правильные ответы.", bold=True)
    _distribution(doc, header.semester, numbered)


def build(parts: Iterable[tuple[FundHeader, list[Section], str]], key: bool) -> bytes:
    """Документ из одной или нескольких частей (например, вариантов теста); каждая – с
    новой страницы: (шапка, разделы по компетенциям, подзаголовок)."""
    doc = _new_document()
    for i, (header, sections, title) in enumerate(parts):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _fund(doc, header, sections, key, title)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- «роза ветров» ----------------------------------------------------------------


def chart_document(title: str, image_png: bytes, legend: list[str], rows: list[list[str]]) -> bytes:
    """Диаграмма для отчёта: заголовок, картинка, подпись и таблица значений по осям."""
    doc = _new_document()
    _para(doc, title, bold=True, center=True)
    doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].add_run().add_picture(io.BytesIO(image_png), width=Mm(140))
    for line in legend:
        _para(doc, line, center=True)
    if rows:
        _para(doc)
        table = _table(doc, len(rows), [25, 117, 25])
        for row, values in zip(table.rows, rows, strict=True):
            for cell, value in zip(row.cells, [*values, "", ""][:3], strict=True):
                _write(cell, value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
